#!/usr/bin/env python3
import tkinter as tk
from tkinter import filedialog
import tkinter.font as tkfont
import json, os, sys, io, base64, copy, uuid, time, webbrowser, re, colorsys, math, traceback
from pathlib import Path

try:
    from PIL import Image, ImageTk
except ImportError:
    print("Pillow is required:  pip install pillow")
    sys.exit(1)
try:
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.colors import HexColor
    from reportlab.lib.utils import ImageReader
except ImportError:
    print("reportlab is required:  pip install reportlab")
    sys.exit(1)
try:
    import cv2
except ImportError:
    cv2 = None

APP_NAME = "PDFS"
HOME_URL = "https://www.hexa-mix.com"
if sys.platform.startswith("win"):
    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("hexamix.pdfs")
    except Exception:
        pass
SETTINGS_FILE = os.path.join(os.path.expanduser("~"), ".pdfstudio.json")

C = {"bg": "#1e1e1e", "panel": "#2a2a2a", "panel2": "#242424", "panel3": "#303030", "border": "#3a3a3a",
     "border2": "#4a4a4a", "topbar": "#0d0d0d", "accent": "#e8a020", "accent2": "#b84820", "ok": "#3a8850",
     "warn": "#c8781a", "err": "#b03030", "text": "#d0d0d0", "text2": "#888888", "text3": "#585858",
     "btn": "#2e2e2e", "btn_hover": "#3e3e3e", "inset": "#0e0e0e", "view": "#161616", "guide": "#ff3fa4"}
THEMES = {
    "HexaMix Dark": {},
    "Midnight": {"bg": "#14181f", "panel": "#1d232d", "panel2": "#181d26", "panel3": "#252c38", "border": "#2c3544",
                 "border2": "#3c485c", "topbar": "#0a0d12", "accent": "#4aa8e8", "accent2": "#2f6fb8", "btn": "#222a36",
                 "btn_hover": "#2e3848", "inset": "#0b0e13", "view": "#10141a", "text": "#d4dae4", "text2": "#8795ab",
                 "text3": "#55627a"},
    "Graphite": {"bg": "#262626", "panel": "#323232", "panel2": "#2c2c2c", "panel3": "#3a3a3a", "border": "#454545",
                 "border2": "#585858", "topbar": "#161616", "accent": "#7ec850", "accent2": "#4a8a2a", "btn": "#383838",
                 "btn_hover": "#484848", "inset": "#181818", "view": "#1c1c1c"},
}
FONT = ("Consolas", 9)
FONTB = ("Consolas", 9, "bold")
SIZES = {"16:9  (960×540)": (960, 540), "4:3  (960×720)": (960, 720), "A4 Landscape  (842×595)": (842, 595),
         "A4 Portrait  (595×842)": (595, 842), "Square  (720×720)": (720, 720)}
TYPE_NAMES = {"text": "Text", "rect": "Rectangle", "ellipse": "Ellipse", "line": "Line", "button": "Button",
              "image": "Image", "video": "Video"}
TYPE_ICON = {"text": "T", "rect": "▭", "ellipse": "◯", "line": "╱", "button": "▣", "image": "▨", "video": "▶"}
TEXTY = ("text", "rect", "ellipse", "button")

FAMILIES = {
    "Arial": ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
    "Times New Roman": ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"),
    "Courier New": ("cour.ttf", "courbd.ttf", "couri.ttf", "courbi.ttf"),
    "Calibri": ("calibri.ttf", "calibrib.ttf", "calibrii.ttf", "calibriz.ttf"),
    "Segoe UI": ("segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf", "segoeuiz.ttf"),
    "Verdana": ("verdana.ttf", "verdanab.ttf", "verdanai.ttf", "verdanaz.ttf"),
    "Georgia": ("georgia.ttf", "georgiab.ttf", "georgiai.ttf", "georgiaz.ttf"),
    "Tahoma": ("tahoma.ttf", "tahomabd.ttf", "tahoma.ttf", "tahomabd.ttf"),
    "Trebuchet MS": ("trebuc.ttf", "trebucbd.ttf", "trebucit.ttf", "trebucbi.ttf"),
    "Consolas": ("consola.ttf", "consolab.ttf", "consolai.ttf", "consolaz.ttf"),
    "Impact": ("impact.ttf",) * 4,
    "DejaVu Sans": ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf"),
    "DejaVu Serif": ("DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf", "DejaVuSerif-BoldItalic.ttf"),
    "Liberation Sans": ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf",
                        "LiberationSans-BoldItalic.ttf"),
    "Liberation Serif": ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf",
                         "LiberationSerif-BoldItalic.ttf"),
}
_font_index = None


def font_index():
    global _font_index
    if _font_index is None:
        _font_index = {}
        dirs = ["C:/Windows/Fonts", "/usr/share/fonts", "/usr/local/share/fonts", "/Library/Fonts",
                "/System/Library/Fonts/Supplemental", os.path.expanduser("~/.fonts"),
                os.path.expanduser("~/Library/Fonts"), os.path.expanduser("~/.local/share/fonts")]
        for d in dirs:
            if os.path.isdir(d):
                for root, _, files in os.walk(d):
                    for f in files:
                        if f.lower().endswith(".ttf"):
                            _font_index.setdefault(f.lower(), os.path.join(root, f))
    return _font_index


def font_path(fam, b, i):
    names = FAMILIES.get(fam)
    if not names:
        return None
    idx = (1 if b else 0) + (2 if i else 0)
    ix = font_index()
    return ix.get(names[idx].lower()) or ix.get(names[0].lower())


def available_families():
    out = [f for f in FAMILIES if font_path(f, False, False)]
    return out or ["Helvetica"]


_registered = {}


def pdf_font(fam, b, i):
    key = f"{fam}|{int(b)}{int(i)}"
    if key in _registered:
        return _registered[key]
    p = font_path(fam, b, i)
    name = None
    if p:
        try:
            pdfmetrics.registerFont(TTFont(key, p))
            name = key
        except Exception:
            name = None
    if name is None:
        name = {(0, 0): "Helvetica", (1, 0): "Helvetica-Bold", (0, 1): "Helvetica-Oblique",
                (1, 1): "Helvetica-BoldOblique"}[(int(b), int(i))]
    _registered[key] = name
    return name


def wrap_text(text, fname, size, maxw):
    lines = []
    for para in str(text).split("\n"):
        words, cur = para.split(" "), ""
        for w in words:
            t = (cur + " " + w) if cur else w
            if pdfmetrics.stringWidth(t, fname, size) <= maxw or not cur:
                cur = t
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def text_layout(el):
    fname = pdf_font(el["font"], el["b"], el["i"])
    size, pad = el["size"], 6
    lines = wrap_text(el["text"], fname, size, max(10, abs(el["w"]) - 2 * pad))
    lh = size * 1.2
    n = len(lines)
    va = el.get("va", "auto")
    if va == "auto":
        va = "top" if el["t"] == "text" else "middle"
    if va == "top":
        bt = pad
    elif va == "middle":
        bt = (el["h"] - n * lh) / 2
    else:
        bt = el["h"] - pad - n * lh
    return lines, lh, bt


def blend(fg, bg, a):
    try:
        f = [int(fg[i:i + 2], 16) for i in (1, 3, 5)]
        b = [int(bg[i:i + 2], 16) for i in (1, 3, 5)]
        return "#%02x%02x%02x" % tuple(int(f[k] * a + b[k] * (1 - a)) for k in range(3))
    except Exception:
        return fg


def lerp_hex(a, b, t):
    return blend(b, a, t)


def hex_to_rgb_t(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def fmt_num(v, dec=1):
    s = f"{v:.{dec}f}"
    return s.rstrip("0").rstrip(".") if "." in s else s


def new_el(kind, fam, **kw):
    e = dict(t=kind, x=100.0, y=100.0, w=300.0, h=120.0, text="", font=fam, size=28, b=False, i=False,
             color="#222222", align="left", va="auto", fill="", line="", lw=1.0, r=0.0, op=1.0, link="",
             asset="", video="", name="")
    if kind == "text":
        e.update(text="Your text here", w=420, h=80, size=36)
    elif kind == "rect":
        e.update(fill="#e8a020", align="center", color="#111111")
    elif kind == "ellipse":
        e.update(fill="#4a7ec8", align="center", color="#ffffff", w=240, h=240)
    elif kind == "line":
        e.update(line="#e8a020", lw=4.0, h=0.0, w=320)
    elif kind == "button":
        e.update(fill="#e8a020", color="#111111", align="center", r=14, text="Button", w=200, h=56, size=22,
                 b=True, link="https://")
    elif kind == "video":
        e.update(w=480, h=270)
    e.update(kw)
    return e


def make_app_icon(size=256):
    from PIL import ImageDraw
    S = 4 * size
    im = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((0, 0, S - 1, S - 1), radius=int(S * 0.2), fill="#0d0d0d")
    k = S / 22.0

    def poly(pts, col):
        d.polygon([((pts[i] + 1) * k, (pts[i + 1] + 1) * k) for i in range(0, len(pts), 2)], fill=col)
    poly((10, 1, 19, 6, 10, 11, 1, 6), "#e8a020")
    poly((1, 6, 10, 11, 10, 19, 1, 14), "#b84820")
    poly((19, 6, 10, 11, 10, 19, 19, 14), "#7a3414")
    return im.resize((size, size), Image.LANCZOS)


class ImgCache:

    def __init__(self, limit=60):
        self.d, self.limit = {}, limit

    def trim(self):
        if len(self.d) > self.limit:
            self.d.clear()

    def get(self, key, pil, w, h):
        w, h = int(w), int(h)
        if pil is None or w < 1 or h < 1 or w > 6000 or h > 6000:
            return None
        k = (key, w, h)
        if k not in self.d:
            self.d[k] = ImageTk.PhotoImage(pil.resize((w, h), Image.BILINEAR))
        return self.d[k]


class DarkScrollbar(tk.Canvas):
    def __init__(self, parent, orient="vertical", command=None, width=8, bg=None, **kw):
        self.orient = orient
        bg = bg if bg is not None else C["inset"]
        super().__init__(parent, bg=bg, highlightthickness=0,
                         width=(width if orient == "vertical" else 0),
                         height=(width if orient == "horizontal" else 0), **kw)
        self.command = command
        self._lo, self._hi, self._drag, self._hover = 0.0, 1.0, None, False
        self.bind("<Configure>", lambda e: self._redraw())
        self.bind("<Button-1>", self._on_click)
        self.bind("<B1-Motion>", self._on_drag)
        self.bind("<ButtonRelease-1>", lambda e: setattr(self, "_drag", None))
        self.bind("<Enter>", lambda e: self._redraw(True))
        self.bind("<Leave>", lambda e: self._redraw(False))

    def set(self, lo, hi):
        self._lo, self._hi = max(0.0, float(lo)), min(1.0, float(hi))
        self._redraw()

    def _rect(self):
        w, h = max(1, self.winfo_width()), max(1, self.winfo_height())
        if self.orient == "vertical":
            y0, y1 = self._lo * h, self._hi * h
            if y1 - y0 < 16:
                y1 = y0 + 16
            return 1, y0, max(1, w - 1), min(h, y1)
        x0, x1 = self._lo * w, self._hi * w
        if x1 - x0 < 16:
            x1 = x0 + 16
        return x0, 1, min(w, x1), max(1, h - 1)

    def _redraw(self, hover=None):
        if hover is not None:
            self._hover = hover
        self.delete("all")
        if self._lo <= 0.0 and self._hi >= 1.0:
            return
        x0, y0, x1, y1 = self._rect()
        r = min(4, (x1 - x0) / 2, (y1 - y0) / 2)
        col = C["border2"] if self._hover else C["panel3"]
        pts = [x0 + r, y0, x1 - r, y0, x1, y0, x1, y0 + r, x1, y1 - r, x1, y1, x1 - r, y1,
               x0 + r, y1, x0, y1, x0, y1 - r, x0, y0 + r, x0, y0]
        self.create_polygon(pts, smooth=True, fill=col)

    def _on_click(self, e):
        x0, y0, x1, y1 = self._rect()
        pos = e.x if self.orient == "horizontal" else e.y
        t0, t1 = (x0, x1) if self.orient == "horizontal" else (y0, y1)
        if t0 <= pos <= t1:
            self._drag = (pos, self._lo)
        elif self.command:
            self.command("scroll", -1 if pos < t0 else 1, "pages")

    def _on_drag(self, e):
        if self._drag is None or not self.command:
            return
        length = self.winfo_width() if self.orient == "horizontal" else self.winfo_height()
        if length <= 0:
            return
        pos = e.x if self.orient == "horizontal" else e.y
        lo = max(0.0, min(1.0 - (self._hi - self._lo), self._drag[1] + (pos - self._drag[0]) / length))
        self.command("moveto", lo)


class DarkCheck(tk.Frame):
    def __init__(self, parent, text, variable, command=None, bg=None, fg=None, font=("Consolas", 9, "bold")):
        bg = bg or C["panel"]
        super().__init__(parent, bg=bg)
        self.variable, self.command = variable, command
        self._box = tk.Canvas(self, width=14, height=14, bg=bg, highlightthickness=0, cursor="hand2")
        self._box.pack(side="left", padx=(2, 7), pady=2)
        self._lbl = tk.Label(self, text=text, bg=bg, fg=fg or C["text"], font=font, cursor="hand2", anchor="w")
        self._lbl.pack(side="left", fill="x", expand=True)
        for w in (self._box, self._lbl):
            w.bind("<Button-1>", self._click)
        self.variable.trace_add("write", lambda *a: self._redraw())
        self._redraw()

    def _click(self, e=None):
        self.variable.set(0 if self.variable.get() else 1)
        if self.command:
            self.command()

    def _redraw(self):
        c = self._box
        try:
            c.delete("all")
            on = bool(self.variable.get())
            c.create_rectangle(1, 1, 13, 13, fill=C["inset"], outline=C["accent"] if on else C["border2"])
            if on:
                c.create_line(3, 7, 6, 10, 11, 3, fill=C["accent"], width=2, joinstyle="round", capstyle="round")
        except tk.TclError:
            pass


class Tip:
    def __init__(self, widget, text):
        self.w, self.text, self.top, self._job = widget, text, None, None
        widget.bind("<Enter>", self._enter, add="+")
        widget.bind("<Leave>", self._leave, add="+")
        widget.bind("<ButtonPress>", self._leave, add="+")

    def _enter(self, e=None):
        self._leave()
        self._job = self.w.after(500, self._show)

    def _leave(self, e=None):
        if self._job:
            try:
                self.w.after_cancel(self._job)
            except Exception:
                pass
            self._job = None
        if self.top:
            try:
                self.top.destroy()
            except Exception:
                pass
            self.top = None

    def _show(self):
        self._job = None
        try:
            x = self.w.winfo_rootx() + 8
            y = self.w.winfo_rooty() + self.w.winfo_height() + 4
            t = tk.Toplevel(self.w)
            t.overrideredirect(True)
            t.configure(bg=C["border2"])
            tk.Label(t, text=self.text, bg=C["panel3"], fg=C["text"], font=("Consolas", 8), padx=8, pady=3,
                     justify="left").pack(padx=1, pady=1)
            t.geometry(f"+{x}+{y}")
            self.top = t
        except Exception:
            self.top = None


def tip(widget, text):
    if text:
        Tip(widget, text)
    return widget


class CtxMenu:
    _open = []
    _t0 = 0.0
    _owner = None
    _away_id = None

    def __init__(self, app):
        self.app = app
        self.items = []
        self.top = None
        self.sub = None

    def add_command(self, label="", command=None, accel="", check=None):
        if not accel and re.search(r"\s{3,}", label):
            label, accel = re.split(r"\s{3,}", label, maxsplit=1)
        self.items.append(("cmd", label, command, accel, check))

    def add_separator(self):
        self.items.append(("sep",))

    def add_cascade(self, label="", menu=None):
        self.items.append(("sub", label, menu))

    def popup(self, x, y, as_sub=False):
        app = self.app
        if not as_sub:
            CtxMenu.close_all()
        t = tk.Toplevel(app.root)
        t.overrideredirect(True)
        t.configure(bg=C["panel3"], highlightthickness=1, highlightbackground=C["border2"])
        self.top = t
        for it in self.items:
            if it[0] == "sep":
                tk.Frame(t, bg=C["border"], height=1).pack(fill="x", padx=3, pady=2)
                continue
            f = tk.Frame(t, bg=C["panel3"])
            f.pack(fill="x")
            if it[0] == "cmd":
                mark = "✓" if it[4] else (" " if it[4] is not None else "")
                if it[4] is not None:
                    tk.Label(f, text=mark, bg=C["panel3"], fg=C["accent"], font=("Consolas", 9, "bold"), padx=6).pack(side="left")
            lb = tk.Label(f, text=it[1], bg=C["panel3"], fg=C["text"], font=("Consolas", 9), padx=12, pady=4, anchor="w")
            lb.pack(side="left", fill="x", expand=True)
            if it[0] == "cmd":
                if it[3]:
                    tk.Label(f, text=it[3], bg=C["panel3"], fg=C["text3"], font=("Consolas", 8), padx=10).pack(side="right")
                for w in [f, lb] + list(f.winfo_children()):
                    w.bind("<Button-1>", lambda e, c=it[2]: self._run(c))
                    w.bind("<Enter>", lambda e, fw=f: self._hover(fw, None))
                    w.bind("<Leave>", lambda e, fw=f: self._unhover(fw))
            else:
                tk.Label(f, text="▸", bg=C["panel3"], fg=C["accent"], font=("Consolas", 9), padx=10).pack(side="right")
                for w in [f, lb] + list(f.winfo_children()):
                    w.bind("<Button-1>", lambda e, fw=f, m=it[2]: self._hover(fw, m))
                    w.bind("<Enter>", lambda e, fw=f, m=it[2]: self._hover(fw, m))
                    w.bind("<Leave>", lambda e, fw=f: self._unhover(fw))
        t.update_idletasks()
        W, H = t.winfo_reqwidth(), t.winfo_reqheight()
        sw, sh = t.winfo_screenwidth(), t.winfo_screenheight()
        x = max(0, min(x, sw - W - 2))
        y = max(0, min(y, sh - H - 2))
        t.geometry(f"{W}x{H}+{x}+{y}")
        t.lift()
        CtxMenu._open.append(self)
        CtxMenu._t0 = time.time()
        if CtxMenu._away_id is None:
            CtxMenu._away_id = app.root.bind_all("<ButtonPress>", CtxMenu._away, add="+")
            app.root.bind_all("<Escape>", lambda e: CtxMenu.close_all(), add="+")

    tk_popup = popup

    def _hover(self, f, sub):
        for ch in f.master.winfo_children():
            if isinstance(ch, tk.Frame):
                self._tint(ch, C["panel3"])
        self._tint(f, C["panel2"])
        self._close_sub()
        if sub is not None:
            self.sub = sub
            sub.popup(f.winfo_rootx() + f.winfo_width() - 2, f.winfo_rooty() - 2, as_sub=True)

    def _unhover(self, f):
        if self.sub is None:
            self._tint(f, C["panel3"])

    @staticmethod
    def _tint(f, col):
        try:
            f.config(bg=col)
            for ch in f.winfo_children():
                ch.config(bg=col)
        except tk.TclError:
            pass

    def _close_sub(self):
        if self.sub is not None:
            self.sub._destroy()
            self.sub = None

    def _destroy(self):
        self._close_sub()
        if self in CtxMenu._open:
            CtxMenu._open.remove(self)
        if self.top is not None:
            try:
                self.top.destroy()
            except Exception:
                pass
            self.top = None

    def _run(self, cmd):
        CtxMenu.close_all()
        if cmd:
            self.app.root.after(10, cmd)

    @staticmethod
    def close_all():
        for m in list(CtxMenu._open):
            m._destroy()
        CtxMenu._open.clear()

    @staticmethod
    def _away(e):
        if not CtxMenu._open or time.time() - CtxMenu._t0 < 0.2:
            return
        try:
            top = e.widget.winfo_toplevel()
        except Exception:
            top = None
        if any(m.top is top for m in CtxMenu._open):
            return
        CtxMenu.close_all()


class CustomMessageBox(tk.Toplevel):

    def __init__(self, parent, title, message, kind="info", accent="accent",
                 yes_text="Yes", no_text="No", ok_text="OK", cancel_text="Cancel"):
        super().__init__(parent)
        self.result = None
        self.overrideredirect(True)
        self.configure(bg=C["panel3"], highlightthickness=1, highlightbackground=C["border2"])
        self._d = (0, 0)
        strip = {"err": C["err"], "warn": C["warn"]}.get(accent, C["accent"])
        fg_on = "#000000" if accent in ("accent", "warn") else "#ffffff"
        tb = tk.Frame(self, bg=C["topbar"], height=26)
        tb.pack(fill="x")
        tb.pack_propagate(False)
        tb.bind("<ButtonPress-1>", lambda e: setattr(self, "_d", (e.x_root - self.winfo_x(), e.y_root - self.winfo_y())))
        tb.bind("<B1-Motion>", lambda e: self.geometry(f"+{e.x_root - self._d[0]}+{e.y_root - self._d[1]}"))
        tk.Label(tb, text=title.upper(), bg=C["topbar"], fg=strip, font=("Consolas", 8, "bold"), padx=8).pack(side="left")
        cl = tk.Label(tb, text=" ✕ ", bg="#601010", fg="#ffaaaa", font=("Consolas", 9, "bold"), cursor="hand2")
        cl.pack(side="right", fill="y")
        cl.bind("<Button-1>", lambda e: self._done(None))
        tk.Frame(self, bg=strip, height=3).pack(fill="x")
        body = tk.Frame(self, bg=C["panel3"])
        body.pack(fill="both", expand=True, padx=16, pady=(14, 10))
        msg = tk.Label(body, text=message, bg=C["panel3"], fg=C["text"], font=("Consolas", 9), justify="left",
                       anchor="w", wraplength=360)
        msg.pack(fill="x")
        btns = tk.Frame(self, bg=C["panel3"])
        btns.pack(fill="x", padx=16, pady=(0, 14))

        def mk(text, bg, fg, val):
            b = tk.Label(btns, text=text, bg=bg, fg=fg, font=("Consolas", 8, "bold"), cursor="hand2", padx=14, pady=5)
            b.bind("<Button-1>", lambda e: self._done(val))
            b.pack(side="right", padx=(8, 0))
            return b
        if kind == "confirm":
            mk(yes_text, strip, fg_on, True)
            mk(no_text, C["btn"], C["text"], False)
            self.bind("<Return>", lambda e: self._done(True))
        elif kind == "triple":
            mk(cancel_text, C["btn"], C["text"], None)
            mk(no_text, C["btn"], C["text"], False)
            mk(yes_text, strip, fg_on, True)
            self.bind("<Return>", lambda e: self._done(True))
        else:
            mk(ok_text, strip, fg_on, True)
            self.bind("<Return>", lambda e: self._done(True))
        self.bind("<Escape>", lambda e: self._done(None))
        self.update_idletasks()
        W, H = max(340, msg.winfo_reqwidth() + 32), self.winfo_reqheight()
        try:
            px = parent.winfo_x() + (parent.winfo_width() - W) // 2
            py = parent.winfo_y() + (parent.winfo_height() - H) // 2
        except Exception:
            px = py = 100
        self.geometry(f"{W}x{H}+{max(0, px)}+{max(0, py)}")
        self.grab_set()
        self.focus_set()

    def _done(self, val):
        self.result = val
        try:
            self.grab_release()
        except Exception:
            pass
        self.destroy()


def _run_dialog(parent, dlg):
    parent.wait_window(dlg)
    return dlg.result


def ask_yes_no(parent, title, message, yes="Yes", no="No"):
    return bool(_run_dialog(parent, CustomMessageBox(parent, title, message, "confirm", yes_text=yes, no_text=no)))


def ask_save_discard(parent, title, message):
    return _run_dialog(parent, CustomMessageBox(parent, title, message, "triple", yes_text="Save",
                                                no_text="Don't save", cancel_text="Cancel"))


def show_info(parent, title, message, accent="accent"):
    return _run_dialog(parent, CustomMessageBox(parent, title, message, "info", accent=accent))


class CustomAskString(tk.Toplevel):
    def __init__(self, parent, title, prompt, initial=""):
        super().__init__(parent)
        self.result = None
        self.overrideredirect(True)
        self.configure(bg=C["panel3"], highlightthickness=1, highlightbackground=C["border2"])
        self._d = (0, 0)
        tb = tk.Frame(self, bg=C["topbar"], height=26)
        tb.pack(fill="x")
        tb.pack_propagate(False)
        tb.bind("<ButtonPress-1>", lambda e: setattr(self, "_d", (e.x_root - self.winfo_x(), e.y_root - self.winfo_y())))
        tb.bind("<B1-Motion>", lambda e: self.geometry(f"+{e.x_root - self._d[0]}+{e.y_root - self._d[1]}"))
        tk.Label(tb, text=title.upper(), bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold"), padx=8).pack(side="left")
        cl = tk.Label(tb, text=" ✕ ", bg="#601010", fg="#ffaaaa", font=("Consolas", 9, "bold"), cursor="hand2")
        cl.pack(side="right", fill="y")
        cl.bind("<Button-1>", lambda e: self._cancel())
        body = tk.Frame(self, bg=C["panel3"])
        body.pack(fill="both", expand=True, padx=16, pady=(14, 10))
        tk.Label(body, text=prompt, bg=C["panel3"], fg=C["text"], font=("Consolas", 9), justify="left",
                 anchor="w", wraplength=330).pack(fill="x", pady=(0, 6))
        self._var = tk.StringVar(value=initial)
        en = tk.Entry(body, textvariable=self._var, bg=C["inset"], fg=C["accent"], font=("Consolas", 10, "bold"), relief="flat",
                      insertbackground=C["accent"], highlightthickness=1, highlightbackground=C["border"], highlightcolor=C["accent"])
        en.pack(fill="x", ipady=4)
        en.select_range(0, "end")
        en.bind("<Return>", lambda e: self._ok())
        en.bind("<Escape>", lambda e: self._cancel())
        btns = tk.Frame(self, bg=C["panel3"])
        btns.pack(fill="x", padx=16, pady=(0, 14))
        ok = tk.Label(btns, text="OK", bg=C["accent"], fg="#000000", font=("Consolas", 8, "bold"), cursor="hand2", padx=14, pady=5)
        ok.pack(side="right")
        ok.bind("<Button-1>", lambda e: self._ok())
        cn = tk.Label(btns, text="Cancel", bg=C["btn"], fg=C["text"], font=("Consolas", 8, "bold"), cursor="hand2", padx=14, pady=5)
        cn.pack(side="right", padx=(0, 8))
        cn.bind("<Button-1>", lambda e: self._cancel())
        self.update_idletasks()
        W, H = 360, self.winfo_reqheight()
        try:
            px = parent.winfo_x() + (parent.winfo_width() - W) // 2
            py = parent.winfo_y() + (parent.winfo_height() - H) // 2
        except Exception:
            px = py = 100
        self.geometry(f"{W}x{H}+{max(0, px)}+{max(0, py)}")
        self.grab_set()
        en.focus_set()

    def _ok(self):
        self.result = self._var.get()
        self._close()

    def _cancel(self):
        self.result = None
        self._close()

    def _close(self):
        try:
            self.grab_release()
        except Exception:
            pass
        self.destroy()


def ask_string(parent, title, prompt, initial=""):
    return _run_dialog(parent, CustomAskString(parent, title, prompt, initial))


class CustomColorPicker(tk.Toplevel):
    SQ = 170
    HW = 20
    SWATCHES = ["#000000", "#ffffff", "#808080", "#c0392b", "#e8a020", "#f1d04a", "#3a8850", "#2a9fd6",
                "#4a6fd6", "#9b59b6", "#8b5a2b", "#f5cba7"]

    def __init__(self, parent, initial="#e8a020", title="Color"):
        super().__init__(parent)
        self.result = None
        self.overrideredirect(True)
        self.configure(bg=C["panel3"], highlightthickness=1, highlightbackground=C["border2"])
        try:
            r, g, b = hex_to_rgb_t(initial)
            self._h, self._s, self._v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
        except Exception:
            self._h, self._s, self._v = 0.1, 0.9, 0.9
        self._d = (0, 0)
        tb = tk.Frame(self, bg=C["topbar"], height=26)
        tb.pack(fill="x")
        tb.pack_propagate(False)
        tb.bind("<ButtonPress-1>", lambda e: setattr(self, "_d", (e.x_root - self.winfo_x(), e.y_root - self.winfo_y())))
        tb.bind("<B1-Motion>", lambda e: self.geometry(f"+{e.x_root - self._d[0]}+{e.y_root - self._d[1]}"))
        tk.Label(tb, text=title.upper(), bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold"), padx=8).pack(side="left")
        cl = tk.Label(tb, text=" ✕ ", bg="#601010", fg="#ffaaaa", font=("Consolas", 9, "bold"), cursor="hand2")
        cl.pack(side="right", fill="y")
        cl.bind("<Button-1>", lambda e: self._cancel())
        body = tk.Frame(self, bg=C["panel3"])
        body.pack(padx=12, pady=12)
        row = tk.Frame(body, bg=C["panel3"])
        row.pack()
        self.sv = tk.Canvas(row, width=self.SQ, height=self.SQ, highlightthickness=1, highlightbackground=C["border"], cursor="crosshair")
        self.sv.pack(side="left")
        self.sv.bind("<Button-1>", self._sv_click)
        self.sv.bind("<B1-Motion>", self._sv_click)
        tk.Frame(row, bg=C["panel3"], width=8).pack(side="left")
        self.hu = tk.Canvas(row, width=self.HW, height=self.SQ, highlightthickness=1, highlightbackground=C["border"], cursor="sb_v_double_arrow")
        self.hu.pack(side="left")
        self.hu.bind("<Button-1>", self._hue_click)
        self.hu.bind("<B1-Motion>", self._hue_click)
        sw = tk.Frame(body, bg=C["panel3"])
        sw.pack(fill="x", pady=(8, 0))
        for i, col in enumerate(self.SWATCHES):
            b = tk.Label(sw, bg=col, width=2, height=1, cursor="hand2")
            b.grid(row=0, column=i, padx=1)
            b.bind("<Button-1>", lambda e, c=col: self._set_hex(c))
        bt = tk.Frame(body, bg=C["panel3"])
        bt.pack(fill="x", pady=(8, 0))
        self.swatch = tk.Canvas(bt, width=30, height=22, highlightthickness=1, highlightbackground=C["border"])
        self.swatch.pack(side="left")
        tk.Label(bt, text="HEX", bg=C["panel3"], fg=C["text3"], font=("Consolas", 7, "bold")).pack(side="left", padx=(8, 3))
        self.hex = tk.StringVar()
        he = tk.Entry(bt, textvariable=self.hex, bg=C["inset"], fg=C["accent"], font=("Consolas", 9, "bold"), relief="flat", width=9,
                      insertbackground=C["accent"], highlightthickness=1, highlightbackground=C["border"], highlightcolor=C["accent"])
        he.pack(side="left", padx=2)
        he.bind("<Return>", lambda e: (self._set_hex(self.hex.get()), "break")[1])
        ok = tk.Label(bt, text="OK", bg=C["accent"], fg="#000", font=("Consolas", 8, "bold"), cursor="hand2", padx=12, pady=3)
        ok.pack(side="right")
        ok.bind("<Button-1>", lambda e: self._ok())
        cn = tk.Label(bt, text="Cancel", bg=C["btn"], fg=C["text"], font=("Consolas", 8, "bold"), cursor="hand2", padx=10, pady=3)
        cn.pack(side="right", padx=(0, 6))
        cn.bind("<Button-1>", lambda e: self._cancel())
        self._pimg = self._himg = None
        self._draw_hue()
        self._draw_sv()
        self._upd()
        self.bind("<Escape>", lambda e: self._cancel())
        self.update_idletasks()
        px = max(0, parent.winfo_pointerx() - 120)
        py = max(0, parent.winfo_pointery() - 40)
        sw_, sh_ = self.winfo_screenwidth(), self.winfo_screenheight()
        px = min(px, sw_ - self.winfo_reqwidth() - 4)
        py = min(py, sh_ - self.winfo_reqheight() - 4)
        self.geometry(f"+{max(0, px)}+{max(0, py)}")
        self.grab_set()
        self.focus_set()

    def _draw_hue(self):
        S = self.SQ
        im = Image.new("RGB", (self.HW, S))
        px = im.load()
        for y in range(S):
            r, g, b = colorsys.hsv_to_rgb(y / (S - 1), 1, 1)
            col = (int(r * 255), int(g * 255), int(b * 255))
            for x in range(self.HW):
                px[x, y] = col
        self._himg = ImageTk.PhotoImage(im)
        self.hu.delete("all")
        self.hu.create_image(0, 0, image=self._himg, anchor="nw")
        y = self._h * (S - 1)
        self.hu.create_line(0, y, self.HW, y, fill="#ffffff", width=2)
        self.hu.create_line(0, y + 1, self.HW, y + 1, fill="#000000")

    def _draw_sv(self):
        S = self.SQ
        hr, hg, hb = colorsys.hsv_to_rgb(self._h, 1, 1)
        data = []
        for y in range(S):
            v = 1 - y / (S - 1)
            for x in range(S):
                s = x / (S - 1)
                data.append((int(((1 - s) + s * hr) * v * 255), int(((1 - s) + s * hg) * v * 255),
                             int(((1 - s) + s * hb) * v * 255)))
        im = Image.new("RGB", (S, S))
        im.putdata(data)
        self._pimg = ImageTk.PhotoImage(im)
        self.sv.delete("all")
        self.sv.create_image(0, 0, image=self._pimg, anchor="nw")
        mx, my = self._s * (S - 1), (1 - self._v) * (S - 1)
        ring = "#ffffff" if self._v < 0.55 else "#000000"
        self.sv.create_oval(mx - 5, my - 5, mx + 5, my + 5, outline=ring, width=2)

    def _cur(self):
        r, g, b = colorsys.hsv_to_rgb(self._h, self._s, self._v)
        return "#%02x%02x%02x" % (round(r * 255), round(g * 255), round(b * 255))

    def _upd(self):
        h = self._cur()
        self.swatch.delete("all")
        self.swatch.create_rectangle(0, 0, 32, 24, fill=h, outline="")
        self.hex.set(h)

    def _sv_click(self, e):
        x, y = max(0, min(self.SQ - 1, e.x)), max(0, min(self.SQ - 1, e.y))
        self._s, self._v = x / (self.SQ - 1), 1 - y / (self.SQ - 1)
        self._draw_sv()
        self._upd()

    def _hue_click(self, e):
        self._h = max(0, min(self.SQ - 1, e.y)) / (self.SQ - 1)
        self._draw_hue()
        self._draw_sv()
        self._upd()

    def _set_hex(self, h):
        h = h.strip()
        if not h.startswith("#"):
            h = "#" + h
        try:
            if len(h) == 4:
                h = "#" + "".join(ch * 2 for ch in h[1:])
            r, g, b = hex_to_rgb_t(h)
            self._h, self._s, self._v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
        except Exception:
            return
        self._draw_hue()
        self._draw_sv()
        self._upd()

    def _ok(self):
        self.result = self._cur()
        self._close()

    def _cancel(self):
        self.result = None
        self._close()

    def _close(self):
        try:
            self.grab_release()
        except Exception:
            pass
        self.destroy()


def pick_color(parent, initial="#e8a020", title="Color"):
    return _run_dialog(parent, CustomColorPicker(parent, initial or "#ffffff", title))


class ToolWindow(tk.Toplevel):

    def __init__(self, app, title, w, h, icon="▣"):
        super().__init__(app.root)
        self.app = app
        self.configure(bg=C["bg"])
        self.overrideredirect(True)
        px = app.root.winfo_x() + max(40, (app.root.winfo_width() - w) // 2)
        py = app.root.winfo_y() + max(30, (app.root.winfo_height() - h) // 2)
        self.geometry(f"{w}x{h}+{px}+{py}")
        self._d = (0, 0)
        tb = tk.Frame(self, bg=C["topbar"], height=28)
        tb.pack(fill="x")
        tb.pack_propagate(False)
        tb.bind("<ButtonPress-1>", self._tb_press)
        tb.bind("<B1-Motion>", self._tb_drag)
        l = tk.Label(tb, text=f"{icon}  {title}", bg=C["topbar"], fg=C["accent"], font=("Consolas", 9, "bold"), padx=8)
        l.pack(side="left")
        l.bind("<ButtonPress-1>", self._tb_press)
        l.bind("<B1-Motion>", self._tb_drag)
        x = tk.Label(tb, text=" ✕ ", bg="#601010", fg="#ffaaaa", font=("Consolas", 9, "bold"), cursor="hand2")
        x.pack(side="right", fill="y")
        x.bind("<Button-1>", lambda e: self.close())
        x.bind("<Enter>", lambda e: x.config(bg="#902020"))
        x.bind("<Leave>", lambda e: x.config(bg="#601010"))
        self.body = tk.Frame(self, bg=C["bg"], highlightthickness=1, highlightbackground=C["border"])
        self.body.pack(fill="both", expand=True)
        self.bind("<Escape>", lambda e: self.close())

    def _tb_press(self, e):
        self._d = (e.x_root - self.winfo_x(), e.y_root - self.winfo_y())

    def _tb_drag(self, e):
        self.geometry(f"+{e.x_root - self._d[0]}+{e.y_root - self._d[1]}")

    def close(self):
        self.destroy()
        try:
            self.app.cv.focus_force()
        except Exception:
            pass


class ScrollFrame(tk.Frame):

    def __init__(self, parent, bg=None, width=None):
        bg = bg or C["panel2"]
        super().__init__(parent, bg=bg)
        self.cv = tk.Canvas(self, bg=bg, highlightthickness=0, bd=0, width=width or 10)
        self.sb = DarkScrollbar(self, orient="vertical", command=self.cv.yview, width=8, bg=bg)
        self.cv.pack(side="left", fill="both", expand=True)
        self.sb.pack(side="right", fill="y")
        self.cv.configure(yscrollcommand=self.sb.set)
        self.inner = tk.Frame(self.cv, bg=bg)
        self._win = self.cv.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda e: self.cv.configure(scrollregion=self.cv.bbox("all")))
        self.cv.bind("<Configure>", lambda e: self.cv.itemconfig(self._win, width=e.width))
        self.bind_wheel(self.cv)
        self.bind_wheel(self.inner)

    def _wheel(self, e):
        if self.inner.winfo_reqheight() <= self.cv.winfo_height():
            return
        self.cv.yview_scroll(-1 if (e.num == 4 or (getattr(e, "delta", 0) > 0)) else 1, "units")

    def bind_wheel(self, widget):
        widget.bind("<MouseWheel>", self._wheel, add="+")
        widget.bind("<Button-4>", self._wheel, add="+")
        widget.bind("<Button-5>", self._wheel, add="+")

    def bind_tree(self, widget=None, skip=()):
        widget = widget or self.inner
        if widget in skip:
            return
        self.bind_wheel(widget)
        for ch in widget.winfo_children():
            self.bind_tree(ch, skip)


def make_section(parent, title, open_=True):
    outer = tk.Frame(parent, bg=C["panel"], highlightthickness=1, highlightbackground=C["border"])
    hdr = tk.Frame(outer, bg=C["topbar"], cursor="hand2")
    hdr.pack(fill="x")
    arrow = tk.Label(hdr, text="▾" if open_ else "▸", bg=C["topbar"], fg=C["accent"],
                     font=("Consolas", 8, "bold"), padx=6, pady=3)
    arrow.pack(side="left")
    lab = tk.Label(hdr, text=title.upper(), bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold"), pady=3)
    lab.pack(side="left")
    body = tk.Frame(outer, bg=C["panel"])
    if open_:
        body.pack(fill="x", padx=6, pady=(4, 5))

    def toggle(e=None):
        if body.winfo_manager():
            body.pack_forget()
            arrow.config(text="▸")
        else:
            body.pack(fill="x", padx=6, pady=(4, 5))
            arrow.config(text="▾")
    for w in (hdr, arrow, lab):
        w.bind("<Button-1>", toggle)
    return body


def draw_icon(c, kind, col):
    L, R, O, P, T = c.create_line, c.create_rectangle, c.create_oval, c.create_polygon, c.create_text
    if kind == "text":
        T(15, 14, text="T", fill=col, font=("Georgia", 15, "bold"))
    elif kind == "rect":
        R(7, 8, 23, 21, outline=col, width=2)
    elif kind == "ellipse":
        O(6, 7, 24, 22, outline=col, width=2)
    elif kind == "line":
        L(7, 21, 23, 8, fill=col, width=2, capstyle="round")
    elif kind == "button":
        R(5, 9, 25, 20, outline=col, width=2)
        L(10, 14, 20, 14, fill=col, width=2)
    elif kind == "image":
        R(6, 7, 24, 22, outline=col, width=2)
        P(8, 20, 13, 13, 17, 18, 19, 15, 23, 20, fill=col, outline="")
        O(18, 9, 21, 12, fill=col, outline="")
    elif kind == "video":
        R(5, 8, 19, 21, outline=col, width=2)
        P(20, 12, 25, 9, 25, 20, 20, 17, fill=col, outline="")
    elif kind == "undo":
        T(15, 13, text="↶", fill=col, font=("Segoe UI Symbol", 16))
    elif kind == "redo":
        T(15, 13, text="↷", fill=col, font=("Segoe UI Symbol", 16))
    elif kind == "new":
        P(9, 6, 18, 6, 22, 10, 22, 23, 9, 23, outline=col, fill="", width=2)
        L(18, 6, 18, 10, 22, 10, fill=col, width=2)
    elif kind == "open":
        P(5, 9, 12, 9, 14, 12, 23, 12, 23, 22, 5, 22, outline=col, fill="", width=2)
    elif kind == "save":
        R(7, 6, 23, 23, outline=col, width=2)
        R(10, 6, 20, 12, outline=col, width=1)
        R(10, 16, 20, 23, outline=col, width=1)
    elif kind == "dup":
        R(6, 6, 18, 18, outline=col, width=2)
        R(11, 11, 23, 23, outline=col, width=2)
    elif kind == "trash":
        R(9, 10, 21, 23, outline=col, width=2)
        L(7, 9, 23, 9, fill=col, width=2)
        L(12, 6, 18, 6, fill=col, width=2)
    elif kind == "front":
        R(6, 6, 17, 17, outline=col, width=1)
        R(12, 12, 23, 23, outline=col, fill=col, width=1)
    elif kind == "back":
        R(12, 12, 23, 23, outline=col, width=1)
        R(6, 6, 17, 17, outline=col, fill=col, width=1)
    elif kind == "play":
        P(10, 7, 10, 22, 23, 14.5, fill=col, outline="")
    elif kind == "pdf":
        T(15, 14, text="PDF", fill=col, font=("Consolas", 9, "bold"))
    elif kind in ("al_l", "al_c", "al_r"):
        ws = (14, 9, 14) if kind != "al_c" else (14, 8, 14)
        for i, w in enumerate((16, 10, 14, 8)):
            y = 8 + i * 4
            x0 = 7 if kind == "al_l" else (15 - w / 2 if kind == "al_c" else 23 - w)
            L(x0, y, x0 + w, y, fill=col, width=2)
    elif kind in ("va_t", "va_m", "va_b"):
        R(6, 6, 24, 23, outline=col, width=1)
        y = {"va_t": 8, "va_m": 12, "va_b": 16}[kind]
        R(10, y, 20, y + 4, fill=col, outline="")
    elif kind in ("a_l", "a_c", "a_r", "a_t", "a_m", "a_b"):
        if kind in ("a_l", "a_c", "a_r"):
            x = {"a_l": 7, "a_c": 15, "a_r": 23}[kind]
            L(x, 5, x, 24, fill=col, width=2)
            x0 = x if kind == "a_l" else (x - 6 if kind == "a_c" else x - 12)
            R(x0 + (1 if kind == "a_l" else 0), 8, x0 + 12, 13, fill=col, outline="")
            x1 = x if kind == "a_l" else (x - 4 if kind == "a_c" else x - 8)
            R(x1 + (1 if kind == "a_l" else 0), 16, x1 + 8, 21, fill=col, outline="")
        else:
            y = {"a_t": 7, "a_m": 14, "a_b": 22}[kind]
            L(5, y, 24, y, fill=col, width=2)
            y0 = y if kind == "a_t" else (y - 6 if kind == "a_m" else y - 12)
            R(8, y0 + (1 if kind == "a_t" else 0), 13, y0 + 12, fill=col, outline="")
            y1 = y if kind == "a_t" else (y - 4 if kind == "a_m" else y - 8)
            R(16, y1 + (1 if kind == "a_t" else 0), 21, y1 + 8, fill=col, outline="")
    elif kind == "dist_h":
        L(6, 5, 6, 24, fill=col, width=2)
        L(24, 5, 24, 24, fill=col, width=2)
        R(12, 9, 18, 20, fill=col, outline="")
    elif kind == "dist_v":
        L(5, 6, 24, 6, fill=col, width=2)
        L(5, 23, 24, 23, fill=col, width=2)
        R(9, 11, 20, 18, fill=col, outline="")
    elif kind == "plus":
        L(8, 14, 22, 14, fill=col, width=2)
        L(15, 7, 15, 21, fill=col, width=2)
    elif kind == "minus":
        L(8, 14, 22, 14, fill=col, width=2)


class IconBtn(tk.Canvas):

    def __init__(self, parent, kind, cmd, tip_text=None, size=(30, 28), bg=None, accent_active=True, hot=None):
        self.base_bg = bg or C["panel"]
        if isinstance(size, int):
            size = (size, size)
        super().__init__(parent, width=size[0], height=size[1], bg=self.base_bg, highlightthickness=0, cursor="hand2")
        self.kind, self.cmd, self.active, self.hover, self.hot = kind, cmd, False, False, hot
        self.bind("<Button-1>", lambda e: self.cmd() if self.cmd else None)
        self.bind("<Enter>", lambda e: self._set_hover(True))
        self.bind("<Leave>", lambda e: self._set_hover(False))
        tip(self, tip_text)
        self._draw()

    def _set_hover(self, h):
        self.hover = h
        self._draw()

    def set_active(self, a):
        a = bool(a)
        if a != self.active:
            self.active = a
            self._draw()

    def _draw(self):
        self.delete("all")
        w, h = int(self["width"]), int(self["height"])
        if self.active:
            bg, col = C["accent"], "#000000"
        elif self.hover:
            bg, col = C["btn_hover"], self.hot or C["accent"]
        else:
            bg, col = self.base_bg, C["text"]
        draw_icon(self, self.kind, col)
        dx, dy = (w - 30) / 2, (h - 28) / 2
        if dx or dy:
            self.move("all", dx, dy)
        self.tag_lower(self.create_rectangle(1, 1, w - 1, h - 1, fill=bg, outline=""))


class TextBtn(tk.Label):
    def __init__(self, parent, text, cmd, tip_text=None, bg=None, fg=None, px=10, py=4, size=9, hot_bg=None):
        self.nbg, self.nfg = bg or C["btn"], fg or C["text"]
        super().__init__(parent, text=text, bg=self.nbg, fg=self.nfg, font=("Consolas", size, "bold"), cursor="hand2",
                         padx=px, pady=py)
        self.cmd, self.hot = cmd, hot_bg or C["btn_hover"]
        self.bind("<Button-1>", lambda e: self.cmd() if self.cmd else None)
        self.bind("<Enter>", lambda e: self.config(bg=self.hot))
        self.bind("<Leave>", lambda e: self.config(bg=self.nbg))
        tip(self, tip_text)


class Swatch(tk.Canvas):

    def __init__(self, parent, on_pick, on_clear=None, w=30, h=20, label="", bg=None):
        super().__init__(parent, width=w, height=h, bg=bg or C["panel"], highlightthickness=1,
                         highlightbackground=C["border2"], cursor="hand2")
        self.w, self.h, self.label, self.col = w, h, label, ""
        self.bind("<Button-1>", lambda e: on_pick())
        if on_clear:
            self.bind("<Button-3>", lambda e: on_clear())
        self.set("")

    def set(self, col):
        self.col = col or ""
        self.delete("all")
        if self.col:
            self.create_rectangle(0, 0, self.w, self.h, fill=self.col, outline="")
        else:
            self.create_rectangle(0, 0, self.w, self.h, fill=C["inset"], outline="")
            self.create_line(2, self.h - 2, self.w - 2, 2, fill=C["err"], width=2)
        if self.label:
            lum = 0.5
            if self.col:
                try:
                    r, g, b = hex_to_rgb_t(self.col)
                    lum = (0.299 * r + 0.587 * g + 0.114 * b) / 255
                except Exception:
                    pass
            self.create_text(self.w / 2, self.h / 2 - 1, text=self.label, fill="#000000" if lum > 0.55 else "#ffffff",
                             font=("Consolas", 9, "bold"))


class NumEntry(tk.Frame):

    def __init__(self, parent, cmd, width=7, step=1.0, lo=None, hi=None, dec=1, bg=None):
        super().__init__(parent, bg=bg or C["panel"])
        self.cmd, self.step, self.lo, self.hi, self.dec = cmd, step, lo, hi, dec
        self.var = tk.StringVar()
        self.e = tk.Entry(self, textvariable=self.var, width=width, bg=C["inset"], fg=C["accent"], font=FONTB,
                          relief="flat", insertbackground=C["accent"], highlightthickness=1,
                          highlightbackground=C["border"], highlightcolor=C["accent"], justify="right")
        self.e.pack(side="left", ipady=2)
        self.e.bind("<Return>", self._commit)
        self.e.bind("<KP_Enter>", self._commit)
        self.e.bind("<FocusOut>", self._commit)
        self.e.bind("<Up>", lambda e: self._bump(1, e))
        self.e.bind("<Down>", lambda e: self._bump(-1, e))
        for ev in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            self.e.bind(ev, self._wheel)

    def set(self, v):
        self.var.set("" if v is None else fmt_num(v, self.dec))

    def _parse(self):
        try:
            return float(self.var.get().replace(",", "."))
        except ValueError:
            return None

    def _clamp(self, v):
        if self.lo is not None:
            v = max(self.lo, v)
        if self.hi is not None:
            v = min(self.hi, v)
        return v

    def _commit(self, e=None):
        v = self._parse()
        if v is not None:
            self.cmd(self._clamp(v))
        if e is not None and getattr(e, "keysym", "") in ("Return", "KP_Enter"):
            self.focus_set()

    def _bump(self, d, e):
        v = self._parse()
        if v is None:
            return "break"
        v = self._clamp(v + d * self.step * (10 if (e.state & 1) else 1))
        self.var.set(fmt_num(v, self.dec))
        self.cmd(v)
        return "break"

    def _wheel(self, e):
        up = e.num == 4 or getattr(e, "delta", 0) > 0
        return self._bump(1 if up else -1, e)


class HSlider(tk.Canvas):

    def __init__(self, parent, lo, hi, cmd, width=170, fmt="{:.0f}", bg=None):
        super().__init__(parent, width=width, height=20, bg=bg or C["panel"], highlightthickness=0, cursor="hand2")
        self.lo, self.hi, self.cmd, self.fmt, self.v = lo, hi, cmd, fmt, lo
        self.bind("<Configure>", lambda e: self._draw())
        self.bind("<Button-1>", self._press)
        self.bind("<B1-Motion>", self._press)

    def set(self, v):
        self.v = max(self.lo, min(self.hi, v))
        self._draw()

    def _track(self):
        w = max(40, self.winfo_width())
        return 6, w - 46

    def _draw(self):
        self.delete("all")
        w = max(40, self.winfo_width())
        x0, x1 = self._track()
        t = (self.v - self.lo) / (self.hi - self.lo) if self.hi > self.lo else 0
        x = x0 + (x1 - x0) * t
        self.create_line(x0, 10, x1, 10, fill=C["inset"], width=4, capstyle="round")
        self.create_line(x0, 10, x, 10, fill=C["accent"], width=4, capstyle="round")
        self.create_oval(x - 6, 4, x + 6, 16, fill=C["text"], outline=C["border2"])
        self.create_text(w - 4, 10, anchor="e", text=self.fmt.format(self.v), fill=C["accent"], font=FONTB)

    def _press(self, e):
        x0, x1 = self._track()
        t = max(0.0, min(1.0, (e.x - x0) / max(1, (x1 - x0))))
        self.v = self.lo + (self.hi - self.lo) * t
        self._draw()
        self.cmd(self.v)


FILLY = ("text", "rect", "ellipse", "button")
GRID = 20


class App:
    def __init__(self, run=True):
        self.settings = self._load_settings()
        self.theme_name = self.settings.get("theme", "HexaMix Dark")
        if self.theme_name in THEMES:
            C.update(THEMES[self.theme_name])
        self.root = tk.Tk()
        self.root.title(APP_NAME)
        self.root.configure(bg=C["bg"])
        self._set_icon()
        self.root.report_callback_exception = self._cb_error
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        w0, h0 = min(1440, sw - 40), min(880, sh - 80)
        self.root.geometry(f"{w0}x{h0}+{(sw - w0) // 2}+{max(0, (sh - h0) // 2 - 20)}")
        self.root.minsize(1140, 640)
        self.root.overrideredirect(True)

        self.fams = available_families()
        self.W, self.H = SIZES["16:9  (960×540)"]
        self.pages = [self.blank_page()]
        self.assets = {}
        self.pil_cache = {}
        self.cache_main = ImgCache(60)
        self.cache_thumb = ImgCache(400)
        self.pi = 0
        self.sel = []
        self.path = None
        self.dirty = False
        self.undo, self.redo = [], []
        self.zoom, self.auto_fit, self.panx, self.pany = 1.0, True, 0.0, 0.0
        self.z, self.ox, self.oy = 1.0, 0.0, 0.0
        self.drag = None
        self.guides = []
        self.marquee = None
        self.editor = None
        self.editing = None
        self.clip = None
        self.space = False
        self.snap_on = bool(self.settings.get("snap", True))
        self.grid_on = bool(self.settings.get("grid", False))
        self.syncing = False
        self._toggles = []
        self._co = (None, 0)
        self.fonts = {}
        self.thumbs = []
        self._thumb_job = None
        self._status_tok = 0
        self._max_geo = None
        self._shown = None
        self._layers_sig = None
        self._drag_ox = self._drag_oy = 0
        self.last_dir = self.settings.get("last_dir", os.path.expanduser("~"))

        self._build()
        self.root.bind("<Key>", self.on_key)
        self.root.bind("<KeyRelease>", self.on_key_up)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.title()
        self.refresh_thumbs()
        self.sync_ui()
        self.root.update_idletasks()
        self.redraw()
        if run:
            self.root.after(150, self.cv.focus_set)
            self.root.mainloop()

    def _set_icon(self):
        try:
            big = make_app_icon(256)
            self._icon_imgs = [ImageTk.PhotoImage(big.resize((s, s), Image.LANCZOS)) for s in (16, 32, 48, 64, 128)]
            self.root.iconphoto(True, *self._icon_imgs)
        except Exception:
            pass
        if sys.platform.startswith("win"):
            try:
                import tempfile
                ico = os.path.join(tempfile.gettempdir(), "pdfs_icon.ico")
                make_app_icon(256).save(ico, sizes=[(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
                self.root.iconbitmap(ico)
            except Exception:
                pass

    def _cb_error(self, exc, val, tb):
        traceback.print_exception(exc, val, tb)
        self.status(f"Error: {val}", "err")

    def _load_settings(self):
        try:
            with open(SETTINGS_FILE) as f:
                return json.load(f)
        except Exception:
            return {}

    def _save_settings(self):
        try:
            with open(SETTINGS_FILE, "w") as f:
                json.dump(dict(snap=self.snap_on, grid=self.grid_on, theme=self.theme_name, last_dir=self.last_dir), f)
        except Exception:
            pass

    def blank_page(self):
        return dict(bg="#ffffff", bg2="#cfd8e6", grad=False, bgimg="", els=[])

    @property
    def page(self):
        return self.pages[self.pi]

    @property
    def els(self):
        return self.page["els"]

    def selels(self):
        return [self.els[i] for i in self.sel if 0 <= i < len(self.els)]

    def cur(self):
        if len(self.sel) == 1 and 0 <= self.sel[0] < len(self.els):
            return self.els[self.sel[0]]
        return None

    def snapshot(self):
        return json.dumps(dict(W=self.W, H=self.H, pages=self.pages, pi=self.pi))

    def push_undo(self):
        self.undo.append(self.snapshot())
        if len(self.undo) > 100:
            self.undo.pop(0)
        self.redo.clear()
        self.dirty = True
        self.title()

    def coalesce(self, key):
        k, t = self._co
        now = time.time()
        if k == key and now - t < 1.2:
            self._co = (key, now)
            return
        self._co = (key, now)
        self.push_undo()

    def load_snap(self, s):
        d = json.loads(s)
        self.W, self.H, self.pages = d["W"], d["H"], d["pages"]
        self.pi = min(d.get("pi", 0), len(self.pages) - 1)
        self.sel = []
        self.end_edit(commit=False)
        self.refresh_thumbs()
        self.sync_ui()
        self.redraw()

    def do_undo(self):
        if self.undo:
            self.redo.append(self.snapshot())
            self.load_snap(self.undo.pop())
            self.dirty = True
            self.title()

    def do_redo(self):
        if self.redo:
            self.undo.append(self.snapshot())
            self.load_snap(self.redo.pop())
            self.dirty = True
            self.title()

    def title(self):
        n = os.path.basename(self.path) if self.path else "Untitled"
        self.root.title(f"{APP_NAME} - {n}{' *' if self.dirty else ''}")
        try:
            self.title_lbl.config(text=f"{n}{'  ●' if self.dirty else ''}")
        except Exception:
            pass

    def status(self, msg, kind="ok"):
        col = {"ok": C["ok"], "warn": C["warn"], "err": C["err"]}.get(kind, C["ok"])
        self.st.config(text=msg, fg=col)
        self._status_tok += 1
        tok = self._status_tok
        self.root.after(7000, lambda: self.st.config(text="") if tok == self._status_tok else None)

    def menu_btn(self, parent, text, builder, bg=None, arrow=False, size=9, tip_text=None, width=None, hover_switch=True):
        nbg = bg or C["btn"]
        b = tk.Label(parent, text=text + (" ▾" if arrow else ""), bg=nbg, fg=C["text"], font=("Consolas", size, "bold"),
                     cursor="hand2", padx=10, pady=4, anchor="w" if width else "center")
        if width:
            b.config(width=width)
        st = {"t": 0.0}

        def show():
            m = CtxMenu(self)
            builder(m)
            b.config(bg=C["accent"], fg="#000000")
            m.popup(b.winfo_rootx(), b.winfo_rooty() + b.winfo_height())
            CtxMenu._owner = b
            b.after(60, watch)

        def watch():
            if CtxMenu._open and CtxMenu._owner is b:
                b.after(80, watch)
            else:
                st["t"] = time.time()
                b.config(bg=nbg, fg=C["text"])

        def click(e=None):
            if CtxMenu._open:
                own = CtxMenu._owner
                CtxMenu.close_all()
                st["t"] = time.time()
                if own is b:
                    return
            elif time.time() - st["t"] < 0.25:
                return
            show()

        def enter(e):
            if CtxMenu._open and CtxMenu._owner is not b and hover_switch:
                CtxMenu.close_all()
                show()
            elif not CtxMenu._open:
                b.config(bg=C["btn_hover"])

        def leave(e):
            if not CtxMenu._open or CtxMenu._owner is not b:
                b.config(bg=nbg)
        b.bind("<Button-1>", click)
        b.bind("<Enter>", enter)
        b.bind("<Leave>", leave)
        tip(b, tip_text)
        return b

    def toggle_btn(self, parent, text, getter, cmd, size=8, px=8, py=3, tip_text=None):
        b = tk.Label(parent, text=text, bg=C["btn"], fg=C["text2"], font=("Consolas", size, "bold"), cursor="hand2",
                     padx=px, pady=py)
        b.bind("<Button-1>", lambda e: cmd())
        b.bind("<Enter>", lambda e: b.config(bg=C["accent"] if getter() else C["btn_hover"]))
        b.bind("<Leave>", lambda e: self.refresh_toggles())
        tip(b, tip_text)
        self._toggles.append((b, getter))
        return b

    def refresh_toggles(self):
        for b, g in self._toggles:
            try:
                on = bool(g())
            except Exception:
                on = False
            try:
                b.config(bg=C["accent"] if on else C["btn"], fg="#000000" if on else C["text2"])
            except tk.TclError:
                pass

    def attr_is(self, k, v):
        els = [e for e in self.selels() if e["t"] in TEXTY]
        if not els:
            return False
        e = els[0]
        cur = e.get(k)
        if k == "va" and cur == "auto":
            cur = "top" if e["t"] == "text" else "middle"
        return cur == v

    def small_btn(self, parent, text, cmd, tip_text=None, px=6, py=4, size=8):
        return TextBtn(parent, text, cmd, tip_text, px=px, py=py, size=size)

    def grid_btns(self, parent, items, cols=2, **kw):
        f = tk.Frame(parent, bg=C["panel"])
        f.pack(fill="x", padx=2, pady=2)
        for i, (txt, cmd) in enumerate(items):
            b = self.small_btn(f, txt, cmd, **kw)
            b.grid(row=i // cols, column=i % cols, padx=1, pady=1, sticky="ew")
            b.config(anchor="w")
        for c in range(cols):
            f.columnconfigure(c, weight=1, uniform="g")
        return f

    def lbl(self, parent, text):
        return tk.Label(parent, text=text, bg=C["panel"], fg=C["text3"], font=("Consolas", 7, "bold"), anchor="w")

    def sec(self, parent, title, open_=True):
        body = make_section(parent, title, open_)
        body.master.pack(fill="x", padx=5, pady=(6, 0))
        return body

    def vsep(self, parent, pad=6):
        tk.Frame(parent, bg=C["border"], width=1).pack(side="left", fill="y", padx=pad, pady=5)

    def _build(self):
        r = self.root
        self._build_titlebar()
        self._build_toolbar()
        self._build_fmtbar()
        sb = tk.Frame(r, bg=C["topbar"], height=20)
        sb.pack(fill="x", side="bottom")
        sb.pack_propagate(False)
        self.sel_lbl = tk.Label(sb, text="", bg=C["topbar"], fg=C["text3"], font=("Consolas", 7))
        self.sel_lbl.pack(side="left", padx=10)
        grip = tk.Label(sb, text="◢", bg=C["topbar"], fg=C["text3"], font=("Consolas", 9),
                        cursor="size_nw_se" if sys.platform.startswith("win") else "bottom_right_corner")
        grip.pack(side="right", padx=2)
        grip.bind("<ButtonPress-1>", lambda e: setattr(self, "_gp", (e.x_root, e.y_root, r.winfo_width(), r.winfo_height())))
        grip.bind("<B1-Motion>", self._grip_drag)
        self.st = tk.Label(sb, text="", bg=C["topbar"], fg=C["ok"], font=("Consolas", 8), anchor="e")
        self.st.pack(side="right", padx=10, fill="x", expand=True)

        body = tk.Frame(r, bg=C["bg"])
        body.pack(fill="both", expand=True)
        left = tk.Frame(body, bg=C["panel2"], width=214, highlightthickness=1, highlightbackground=C["border"])
        left.pack(side="left", fill="y")
        left.pack_propagate(False)
        self._build_left(left)
        right = tk.Frame(body, bg=C["panel2"], width=304, highlightthickness=1, highlightbackground=C["border"])
        right.pack(side="right", fill="y")
        right.pack_propagate(False)
        self._build_right(right)
        mid = tk.Frame(body, bg=C["bg"])
        mid.pack(side="left", fill="both", expand=True)
        vp = tk.Frame(mid, bg=C["panel"], highlightthickness=1, highlightbackground=C["border"])
        vp.pack(fill="both", expand=True, padx=3, pady=3)
        vh = tk.Frame(vp, bg=C["topbar"])
        vh.pack(fill="x")
        tk.Label(vh, text="▾", bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold"), padx=6, pady=3).pack(side="left")
        tk.Label(vh, text="PAGE VIEW", bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold")).pack(side="left")
        self.vp_info = tk.Label(vh, text="", bg=C["topbar"], fg=C["text3"], font=("Consolas", 8))
        self.vp_info.pack(side="right", padx=8)
        self.cv = tk.Canvas(vp, bg=C["view"], highlightthickness=0, cursor="arrow", takefocus=1)
        self.cv.pack(fill="both", expand=True)
        cv = self.cv
        cv.bind("<Configure>", lambda e: self.redraw())
        cv.bind("<ButtonPress-1>", self.on_press)
        cv.bind("<B1-Motion>", self.on_drag)
        cv.bind("<ButtonRelease-1>", self.on_release)
        cv.bind("<Double-Button-1>", self.on_dbl)
        cv.bind("<Motion>", self.on_motion)
        cv.bind("<ButtonPress-3>", self.on_rclick)
        cv.bind("<ButtonPress-2>", self.on_pan_start)
        cv.bind("<B2-Motion>", self.on_pan)
        cv.bind("<MouseWheel>", self.on_wheel)
        cv.bind("<Button-4>", self.on_wheel)
        cv.bind("<Button-5>", self.on_wheel)

    def _build_titlebar(self):
        r = self.root
        tb = tk.Frame(r, bg=C["topbar"], height=30)
        tb.pack(fill="x")
        tb.pack_propagate(False)
        logo = tk.Canvas(tb, width=20, height=20, bg=C["topbar"], highlightthickness=0)
        logo.pack(side="left", padx=(8, 3), pady=5)
        logo.create_polygon(10, 1, 19, 6, 10, 11, 1, 6, fill=C["accent"], outline="")
        logo.create_polygon(1, 6, 10, 11, 10, 19, 1, 14, fill=C["accent2"], outline="")
        logo.create_polygon(19, 6, 10, 11, 10, 19, 19, 14, fill="#7a3414", outline="")
        nm = tk.Label(tb, text="PDFS", bg=C["topbar"], fg=C["accent"], font=("Consolas", 10, "bold"))
        nm.pack(side="left")
        l3 = tk.Label(tb, text="  ·  ", bg=C["topbar"], fg=C["text3"], font=("Consolas", 9))
        l3.pack(side="left")
        l4 = tk.Label(tb, text="Build slides, export PDF", bg=C["topbar"], fg=C["text2"], font=("Consolas", 9),
                      cursor="hand2")
        l4.pack(side="left")
        l4.bind("<Button-1>", lambda e: webbrowser.open(HOME_URL))
        l4.bind("<Enter>", lambda e: l4.config(fg=C["accent"], font=("Consolas", 9, "underline")))
        l4.bind("<Leave>", lambda e: l4.config(fg=C["text2"], font=("Consolas", 9)))
        tip(l4, "Open www.hexa-mix.com")
        close = tk.Label(tb, text=" ✕ ", bg="#601010", fg="#ffaaaa", font=("Consolas", 10, "bold"), cursor="hand2")
        close.pack(side="right", fill="y")
        close.bind("<Button-1>", lambda e: self.on_close())
        close.bind("<Enter>", lambda e: close.config(bg="#902020"))
        close.bind("<Leave>", lambda e: close.config(bg="#601010"))
        for txt, cmd in (("  □  ", self._maximize), (" ─ ", self._minimize)):
            l = tk.Label(tb, text=txt, bg=C["topbar"], fg=C["text3"], font=("Consolas", 10, "bold"), cursor="hand2")
            l.pack(side="right", fill="y")
            l.bind("<Button-1>", lambda e, c=cmd: c())
            l.bind("<Enter>", lambda e, l=l: l.config(bg=C["panel"]))
            l.bind("<Leave>", lambda e, l=l: l.config(bg=C["topbar"]))
        self.title_lbl = tk.Label(tb, text="", bg=C["topbar"], fg=C["text2"], font=("Consolas", 8))
        self.title_lbl.pack(side="right", padx=14)
        for w in (tb, logo, nm, l3, self.title_lbl):
            w.bind("<ButtonPress-1>", self._tb_press)
            w.bind("<B1-Motion>", self._tb_drag)
            w.bind("<Double-Button-1>", lambda e: self._maximize())

    def _build_left(self, left):
        s = self.sec(left, "Add Element")
        s.master.pack_configure(padx=3, pady=(4, 0))
        self.grid_btns(s, [("T  Text", lambda: self.add_el("text")), ("▭ Rectangle", lambda: self.add_el("rect")),
                           ("◯ Ellipse", lambda: self.add_el("ellipse")), ("╱ Line", lambda: self.add_el("line")),
                           ("▣ Button", lambda: self.add_el("button")), ("🖼 Image", self.add_image),
                           ("🎬 Video", self.add_video), ("＋ Page", self.add_page)], 2, px=3)
        h = tk.Frame(left, bg=C["topbar"], height=26)
        h.pack(fill="x", padx=3, pady=(8, 0))
        h.pack_propagate(False)
        tk.Label(h, text=" PAGES", bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold")).pack(side="left", padx=4)
        for txt, cmd, tp in (("🗑", self.del_page, "Delete page"), ("⧉", self.dup_page, "Duplicate page"),
                             ("＋", self.add_page, "New page  (Ctrl+M)")):
            b = tk.Label(h, text=f" {txt} ", bg=C["btn"], fg=C["text2"], font=("Consolas", 9), cursor="hand2")
            b.pack(side="right", padx=1, pady=3)
            b.bind("<Button-1>", lambda e, c=cmd: c())
            tip(b, tp)
        self.th = ScrollFrame(left, bg=C["panel2"])
        self.th.pack(fill="both", expand=True, padx=3, pady=(0, 3))

    def _build_right(self, right):
        h = tk.Frame(right, bg=C["topbar"], height=26)
        h.pack(fill="x", padx=3, pady=(4, 0))
        h.pack_propagate(False)
        tk.Label(h, text=" LAYERS", bg=C["topbar"], fg=C["accent"], font=("Consolas", 8, "bold")).pack(side="left", padx=4)
        for txt, cmd, tp in (("🗑", self.del_el, "Delete selected"), ("⧉", self.dup_el, "Duplicate"),
                             ("▼", lambda: self.z_move("back"), "Send to back"), ("▲", lambda: self.z_move("front"), "Bring to front")):
            b = tk.Label(h, text=f" {txt} ", bg=C["btn"], fg=C["text2"], font=("Consolas", 9), cursor="hand2")
            b.pack(side="right", padx=1, pady=3)
            b.bind("<Button-1>", lambda e, c=cmd: c())
            tip(b, tp)
        lf = tk.Frame(right, bg=C["inset"], highlightthickness=1, highlightbackground=C["border"])
        lf.pack(fill="x", padx=3)
        self.layers_sf = ScrollFrame(lf, bg=C["inset"])
        self.layers_sf.cv.config(height=150)
        self.layers_sf.pack(fill="x")
        tabs = tk.Frame(right, bg=C["topbar"])
        tabs.pack(fill="x", padx=3, pady=(6, 0))
        content = tk.Frame(right, bg=C["panel2"])
        content.pack(fill="both", expand=True, padx=3, pady=(0, 3))
        self._rtab_btns, self._rtab_frames = {}, {}
        for tid, label, builder in (("element", "ELEMENT", self._tab_element), ("text", "TEXT", self._tab_text),
                                    ("page", "PAGE", self._tab_page)):
            b = tk.Label(tabs, text=f" {label} ", bg=C["topbar"], fg=C["text3"], font=("Consolas", 8, "bold"), cursor="hand2", pady=5)
            b.pack(side="left")
            b.bind("<Button-1>", lambda e, t=tid: self.set_rtab(t))
            sf = ScrollFrame(content, bg=C["panel2"])
            builder(sf.inner)
            sf.bind_tree(sf.inner)
            self._rtab_btns[tid], self._rtab_frames[tid] = b, sf
        self.set_rtab("element")

    def set_rtab(self, tid):
        for k, b in self._rtab_btns.items():
            b.config(bg=C["panel3"] if k == tid else C["topbar"], fg=C["accent"] if k == tid else C["text3"])
        for k, f in self._rtab_frames.items():
            if k == tid:
                f.pack(fill="both", expand=True)
            else:
                f.pack_forget()

    def _tb_press(self, e):
        self._drag_ox, self._drag_oy = e.x_root - self.root.winfo_x(), e.y_root - self.root.winfo_y()

    def _tb_drag(self, e):
        if self._max_geo:
            return
        self.root.geometry(f"+{e.x_root - self._drag_ox}+{e.y_root - self._drag_oy}")

    def _maximize(self):
        if self._max_geo:
            self.root.geometry(self._max_geo)
            self._max_geo = None
        else:
            self._max_geo = self.root.geometry()
            self.root.geometry(f"{self.root.winfo_screenwidth()}x{self.root.winfo_screenheight()}+0+0")

    def _minimize(self):
        self.root.overrideredirect(False)
        self.root.iconify()

        def restore(e):
            if self.root.winfo_viewable():
                self.root.overrideredirect(True)
                self.root.unbind("<Map>")
        self.root.bind("<Map>", restore)

    def _grip_drag(self, e):
        x0, y0, w, h = self._gp
        self.root.geometry(f"{max(1140, w + e.x_root - x0)}x{max(640, h + e.y_root - y0)}")

    def _m_file(self, m):
        m.add_command(label="New project         Ctrl+N", command=self.new_doc)
        m.add_command(label="Open…               Ctrl+O", command=self.open_doc)
        m.add_command(label="Save                Ctrl+S", command=self.save_doc)
        m.add_command(label="Save As…            Ctrl+Shift+S", command=lambda: self.save_doc(True))
        m.add_separator()
        m.add_command(label="Export PDF…         Ctrl+E", command=self.export_pdf)
        m.add_command(label="Export PowerPoint…  Ctrl+Shift+E", command=self.export_pptx)
        m.add_command(label="Present             F5", command=self.present)
        m.add_separator()
        m.add_command(label="Quit", command=self.on_close)

    def _m_edit(self, m):
        m.add_command(label="Undo                Ctrl+Z", command=self.do_undo)
        m.add_command(label="Redo                Ctrl+Y", command=self.do_redo)
        m.add_separator()
        m.add_command(label="Cut                 Ctrl+X", command=self.cut)
        m.add_command(label="Copy                Ctrl+C", command=self.copy)
        m.add_command(label="Paste               Ctrl+V", command=self.paste)
        m.add_command(label="Duplicate           Ctrl+D", command=self.dup_el)
        m.add_command(label="Delete              Del", command=self.del_el)
        m.add_separator()
        m.add_command(label="Select All          Ctrl+A", command=self.select_all)
        m.add_command(label="Edit Text           F2", command=self.edit_selected)

    def _m_insert(self, m):
        for k in ("text", "rect", "ellipse", "line", "button"):
            m.add_command(label=f"{TYPE_ICON[k]}  {TYPE_NAMES[k]}", command=lambda k=k: self.add_el(k))
        m.add_separator()
        m.add_command(label="🖼  Image…", command=self.add_image)
        m.add_command(label="🎬  Video…", command=self.add_video)
        m.add_separator()
        m.add_command(label="＋  New Page           Ctrl+M", command=self.add_page)

    def _m_arrange(self, m):
        m.add_command(label="Bring to Front", command=lambda: self.z_move("front"))
        m.add_command(label="Bring Forward", command=lambda: self.z_move(1))
        m.add_command(label="Send Backward", command=lambda: self.z_move(-1))
        m.add_command(label="Send to Back", command=lambda: self.z_move("back"))
        m.add_separator()
        am = CtxMenu(self)
        for lab, how in (("Left", "l"), ("Center Horizontally", "c"), ("Right", "r"), ("Top", "t"),
                         ("Center Vertically", "m"), ("Bottom", "b")):
            am.add_command(label=lab, command=lambda h=how: self.align(h))
        m.add_cascade(label="Align", menu=am)
        dm = CtxMenu(self)
        dm.add_command(label="Horizontally", command=lambda: self.distribute("h"))
        dm.add_command(label="Vertically", command=lambda: self.distribute("v"))
        m.add_cascade(label="Distribute", menu=dm)

    def _m_page(self, m):
        m.add_command(label="New Page            Ctrl+M", command=self.add_page)
        m.add_command(label="Duplicate Page", command=self.dup_page)
        m.add_command(label="Delete Page", command=self.del_page)
        m.add_separator()
        m.add_command(label="Move Page Up", command=lambda: self.move_page(-1))
        m.add_command(label="Move Page Down", command=lambda: self.move_page(1))
        m.add_separator()
        sm = CtxMenu(self)
        for n in SIZES:
            sm.add_command(label=n, command=lambda n=n: self.set_size(*SIZES[n]))
        sm.add_separator()
        sm.add_command(label="Custom…", command=self.custom_size)
        m.add_cascade(label="Page Size", menu=sm)

    def _m_view(self, m):
        m.add_command(label="Zoom In               +", command=lambda: self.zoom_by(1.2))
        m.add_command(label="Zoom Out              -", command=lambda: self.zoom_by(1 / 1.2))
        m.add_command(label="Fit to Window         Ctrl+0", command=self.fit)
        m.add_command(label="Actual Size           Ctrl+1", command=lambda: self.set_zoom(1.0))
        m.add_separator()
        m.add_command(label="Snap to Guides", command=self.toggle_snap, check=self.snap_on)
        m.add_command(label="Show Grid", command=self.toggle_grid, check=self.grid_on)
        m.add_separator()
        tm = CtxMenu(self)
        for n in THEMES:
            tm.add_command(label=n, command=lambda n=n: self.set_theme(n), check=(n == self.theme_name))
        m.add_cascade(label="Theme (after restart)", menu=tm)
        m.add_separator()
        m.add_command(label="Keyboard Shortcuts    F1", command=self.show_shortcuts)
        m.add_command(label="About", command=self.show_about)

    def _m_help(self, m):
        m.add_command(label="Keyboard Shortcuts", command=self.show_shortcuts)
        m.add_command(label="About", command=self.show_about)

    def toggle_snap(self):
        self.snap_on = not self.snap_on
        self._save_settings()
        self.status("Snap to guides: " + ("on" if self.snap_on else "off"))

    def toggle_grid(self):
        self.grid_on = not self.grid_on
        self._save_settings()
        self.redraw()

    def set_theme(self, n):
        self.theme_name = n
        self._save_settings()
        show_info(self.root, "Theme", f"Theme \"{n}\" selected.\nIt takes effect on the next start.")

    def _build_toolbar(self):
        bar = tk.Frame(self.root, bg=C["panel"], height=36, highlightthickness=1, highlightbackground=C["border"])
        bar.pack(fill="x")
        bar.pack_propagate(False)
        L = lambda b: b.pack(side="left", padx=2, pady=3)
        L(TextBtn(bar, "＋ New", self.new_doc, "New project  (Ctrl+N)"))
        L(TextBtn(bar, "📂 Open", self.open_doc, "Open project  (Ctrl+O)"))
        L(TextBtn(bar, "💾 Save", self.save_doc, "Save project  (Ctrl+S)  ·  Ctrl+Shift+S = Save As"))
        self.vsep(bar)
        L(TextBtn(bar, "↶", self.do_undo, "Undo  (Ctrl+Z)", px=8))
        L(TextBtn(bar, "↷", self.do_redo, "Redo  (Ctrl+Y)", px=8))
        L(TextBtn(bar, "⧉ Dup", self.dup_el, "Duplicate  (Ctrl+D)"))
        L(TextBtn(bar, "🗑 Del", self.del_el, "Delete  (Del)"))
        self.vsep(bar)
        L(self.menu_btn(bar, "＋ Insert", self._m_insert, arrow=True, tip_text="Insert element  (text, shapes, image, video)"))
        L(self.menu_btn(bar, "Arrange", self._m_arrange, arrow=True, tip_text="Order, align, distribute"))
        L(self.menu_btn(bar, "Page", self._m_page, arrow=True, tip_text="Pages and page size"))
        L(self.menu_btn(bar, "View", self._m_view, arrow=True, tip_text="Zoom, snapping, theme"))
        info = TextBtn(bar, "ⓘ", self.show_shortcuts, "Keyboard shortcuts  (F1)", px=8)
        info.pack(side="right", padx=(2, 6), pady=3)
        TextBtn(bar, "⬇ EXPORT PDF", self.export_pdf, "Export PDF  (Ctrl+E)", bg=C["accent2"], fg="#ffffff",
                hot_bg="#d05a28").pack(side="right", padx=2, pady=3)
        TextBtn(bar, "⬇ PPTX", self.export_pptx, "Export PowerPoint (.pptx)  (Ctrl+Shift+E)", bg="#a8401c", fg="#ffffff",
                hot_bg="#c85a30").pack(side="right", padx=2, pady=3)
        TextBtn(bar, "▶ Present", self.present, "Fullscreen presentation  (F5)", bg="#2a3a52", hot_bg="#34496a").pack(side="right", padx=2, pady=3)

    def _build_fmtbar(self):
        fb = tk.Frame(self.root, bg=C["panel"], height=34, highlightthickness=1, highlightbackground=C["border"])
        fb.pack(fill="x")
        fb.pack_propagate(False)
        small = lambda t: tk.Label(fb, text=t, bg=C["panel"], fg=C["text3"], font=("Consolas", 7, "bold"))
        small("FONT").pack(side="left", padx=(8, 3))
        self.font_btn = self.menu_btn(fb, "—", self._m_fonts, bg=C["inset"], arrow=True, size=8, width=16,
                                      hover_switch=False, tip_text="Font family")
        self.font_btn.pack(side="left", padx=1, pady=3)
        small("SIZE").pack(side="left", padx=(8, 3))
        TextBtn(fb, "−", lambda: self.bump_size(-1), "Smaller text", px=7, py=3, size=8).pack(side="left")
        self.size_ent = NumEntry(fb, lambda v: self.set_attr("size", float(v), texty=True, co="size"), width=4, step=1,
                                 lo=4, hi=400, dec=0, bg=C["panel"])
        self.size_ent.pack(side="left", padx=2)
        TextBtn(fb, "+", lambda: self.bump_size(1), "Larger text", px=7, py=3, size=8).pack(side="left")
        self.vsep(fb)
        tg = lambda t, g, c, tp=None, px=8: self.toggle_btn(fb, t, g, c, size=8, px=px, py=4, tip_text=tp)
        tg("B", lambda: self.attr_is("b", True), lambda: self.toggle("b"), "Bold", 9).pack(side="left", padx=1)
        tg("I", lambda: self.attr_is("i", True), lambda: self.toggle("i"), "Italic", 9).pack(side="left", padx=1)
        self.vsep(fb)
        small("ALIGN").pack(side="left", padx=(2, 3))
        for a, t, tt in (("left", "Left", "Align left"), ("center", "Center", "Align center"), ("right", "Right", "Align right")):
            tg(t, lambda a=a: self.attr_is("align", a), lambda a=a: self.set_attr("align", a, texty=True), tt, 6).pack(side="left", padx=1)
        self.vsep(fb)
        self.swatches = {"color": [], "fill": [], "line": []}
        for key, lab, tt in (("color", "TEXT", "Text color"), ("fill", "FILL", "Fill color  (right-click = none)"),
                             ("line", "LINE", "Border / line color  (right-click = none)")):
            small(lab).pack(side="left", padx=(4, 3))
            sw = Swatch(fb, lambda k=key: self.pick_attr_color(k), (lambda k=key: self.set_color(k, "")) if key != "color" else None,
                        label="A" if key == "color" else "", bg=C["panel"])
            sw.pack(side="left", padx=1)
            tip(sw, tt)
            self.swatches[key].append(sw)
        self.vsep(fb)
        tg("🧲 Snap", lambda: self.snap_on, self.toggle_snap, "Snap to alignment guides while dragging", 7).pack(side="left", padx=1)
        tg("# Grid", lambda: self.grid_on, self.toggle_grid, "Show grid (and snap to it)", 7).pack(side="left", padx=1)
        TextBtn(fb, "100%", lambda: self.set_zoom(1.0), "Actual size  (Ctrl+1)", px=6, py=3, size=8).pack(side="right", padx=(1, 8))
        TextBtn(fb, "⌖ Fit", self.fit, "Fit page to window  (Ctrl+0)", px=6, py=3, size=8).pack(side="right", padx=1)
        TextBtn(fb, "+", lambda: self.zoom_by(1.2), "Zoom in", px=7, py=3, size=8).pack(side="right", padx=1)
        self.zl = tk.Label(fb, text="100%", bg=C["panel"], fg=C["accent"], font=("Consolas", 8, "bold"), width=5)
        self.zl.pack(side="right")
        TextBtn(fb, "−", lambda: self.zoom_by(1 / 1.2), "Zoom out", px=7, py=3, size=8).pack(side="right", padx=1)
        small("ZOOM").pack(side="right", padx=(0, 3))

    def _m_fonts(self, m):
        for f in self.fams:
            m.add_command(label=f, command=lambda f=f: self.set_attr("font", f, texty=True))

    def irow(self, parent, label, w=11):
        f = tk.Frame(parent, bg=C["panel"])
        f.pack(fill="x", pady=2)
        tk.Label(f, text=label, bg=C["panel"], fg=C["text2"], font=("Consolas", 8), width=w, anchor="w").pack(side="left")
        return f

    def _entry(self, parent, var, **kw):
        return tk.Entry(parent, textvariable=var, bg=C["inset"], fg=C["accent"], font=FONTB, relief="flat",
                        insertbackground=C["accent"], highlightthickness=1, highlightbackground=C["border"],
                        highlightcolor=C["accent"], **kw)

    def _tab_element(self, p):
        s = self.sec(p, "Transform")
        self.el_name = tk.Label(s, text="NO SELECTION", bg=C["panel"], fg=C["accent"], font=FONTB, anchor="w")
        self.el_name.pack(fill="x", pady=(0, 3))
        self.lbl(s, "NAME").pack(fill="x")
        self.name_var = tk.StringVar()
        ne = self._entry(s, self.name_var)
        ne.pack(fill="x", ipady=3, pady=(0, 4))
        ne.bind("<Return>", lambda ev: (self.commit_name(), self.cv.focus_set()))
        ne.bind("<FocusOut>", lambda ev: self.commit_name())
        g = tk.Frame(s, bg=C["panel"])
        g.pack(fill="x")
        self.geo = {}
        for i, (k, lab) in enumerate((("x", "X"), ("y", "Y"), ("w", "W"), ("h", "H"))):
            cell = tk.Frame(g, bg=C["panel"])
            cell.grid(row=i // 2, column=i % 2, sticky="ew", padx=2, pady=2)
            tk.Label(cell, text=lab, bg=C["panel"], fg=C["text2"], font=("Consolas", 8, "bold"), width=2, anchor="w").pack(side="left")
            ne2 = NumEntry(cell, lambda v, k=k: self.commit_geom(k, v), width=7, step=1, dec=1)
            ne2.pack(side="left", fill="x", expand=True)
            self.geo[k] = ne2
        g.columnconfigure(0, weight=1)
        g.columnconfigure(1, weight=1)
        f = self.irow(s, "OPACITY", 8)
        self.op_sl = HSlider(f, 0, 100, lambda v: self.set_attr("op", round(v) / 100.0, co="op"), width=150)
        self.op_sl.pack(side="left", fill="x", expand=True)

        s = self.sec(p, "Align")
        self.grid_btns(s, [("Left", lambda: self.align("l")), ("Center", lambda: self.align("c")), ("Right", lambda: self.align("r")),
                           ("Top", lambda: self.align("t")), ("Middle", lambda: self.align("m")), ("Bottom", lambda: self.align("b")),
                           ("Dist. H", lambda: self.distribute("h")), ("Dist. V", lambda: self.distribute("v"))], 3, px=3)
        tk.Label(s, text="One element aligns to the page, several align to each other.", bg=C["panel"], fg=C["text3"],
                 font=("Consolas", 7), anchor="w", justify="left", wraplength=250).pack(fill="x", padx=4, pady=(2, 0))

        s = self.sec(p, "Order")
        self.grid_btns(s, [("⬆ To Front", lambda: self.z_move("front")), ("⬇ To Back", lambda: self.z_move("back")),
                           ("▲ Forward", lambda: self.z_move(1)), ("▼ Backward", lambda: self.z_move(-1)),
                           ("⧉ Duplicate", self.dup_el), ("🗑 Delete", self.del_el)], 2, px=3)

        s = self.sec(p, "Appearance")
        for key, lab in (("fill", "FILL"), ("line", "BORDER")):
            f = self.irow(s, lab, 8)
            clr = tk.Label(f, text="✕", bg=C["inset"], fg=C["text3"], font=FONTB, cursor="hand2", padx=6)
            clr.pack(side="right", padx=(4, 0))
            clr.bind("<Button-1>", lambda e, k=key: self.set_color(k, ""))
            tip(clr, "None")
            sw = Swatch(f, lambda k=key: self.pick_attr_color(k), w=64, h=20)
            sw.pack(side="right")
            self.swatches[key].append(sw)
        f = self.irow(s, "WIDTH", 8)
        self.lw_sl = HSlider(f, 0, 30, lambda v: self.set_attr("lw", round(v * 2) / 2, co="lw"), width=150, fmt="{:.1f}")
        self.lw_sl.pack(side="left", fill="x", expand=True)
        f = self.irow(s, "ROUNDING", 8)
        self.r_sl = HSlider(f, 0, 100, lambda v: self.set_attr("r", round(v), co="r"), width=150)
        self.r_sl.pack(side="left", fill="x", expand=True)

        s = self.sec(p, "Link / Source")
        self.link_var = tk.StringVar()
        e = self._entry(s, self.link_var)
        e.pack(fill="x", ipady=3)
        e.bind("<Return>", lambda ev: (self.commit_link(), self.cv.focus_set()))
        e.bind("<FocusOut>", lambda ev: self.commit_link())
        self.link_hint = tk.Label(s, text="", bg=C["panel"], fg=C["text3"], font=("Consolas", 7), anchor="w", justify="left")
        self.link_hint.pack(fill="x", pady=(2, 3))
        f = tk.Frame(s, bg=C["panel"])
        f.pack(fill="x")
        self.menu_btn(f, "Jump to page", self._m_jump, arrow=True, size=8, hover_switch=False).pack(side="left", padx=(0, 3))
        TextBtn(f, "Clear link", self.clear_link, px=8, py=4, size=8).pack(side="left")
        self.replace_btn = TextBtn(s, "Replace image / poster…", self.replace_asset, px=8, py=4, size=8)
        self.replace_btn.pack(fill="x", pady=(4, 0))

    def _tab_text(self, p):
        s = self.sec(p, "Content")
        self.tx = tk.Text(s, height=7, width=28, bg=C["inset"], fg=C["text"], font=("Consolas", 10),
                          insertbackground=C["accent"], relief="flat", bd=3, highlightthickness=1,
                          highlightbackground=C["border"], highlightcolor=C["accent"], wrap="word", undo=False)
        self.tx.pack(fill="x")
        self.tx.bind("<KeyRelease>", self.commit_text)
        tk.Label(s, text="Double-click an element to edit its text in place.", bg=C["panel"], fg=C["text3"],
                 font=("Consolas", 7), anchor="w").pack(fill="x", pady=(3, 0))
        s = self.sec(p, "Style")
        f = self.irow(s, "COLOR", 8)
        sw = Swatch(f, lambda: self.pick_attr_color("color"), w=64, h=20)
        sw.pack(side="left")
        self.swatches["color"].append(sw)
        f = self.irow(s, "FONT", 8)
        self.menu_btn(f, "Font family", self._m_fonts, arrow=True, size=8, hover_switch=False).pack(side="left")
        f = self.irow(s, "V-ALIGN", 8)
        for a, t in (("top", "Top"), ("middle", "Middle"), ("bottom", "Bottom")):
            self.toggle_btn(f, t, lambda a=a: self.attr_is("va", a), lambda a=a: self.set_attr("va", a, texty=True),
                            size=8, px=8, py=3, tip_text=f"Vertical alignment: {t.lower()}").pack(side="left", padx=1)

    def _tab_page(self, p):
        s = self.sec(p, "Page Size")
        self.size_btn = self.menu_btn(s, "16:9", self._m_sizes, arrow=True, size=8, hover_switch=False)
        self.size_btn.pack(fill="x", pady=2)
        s = self.sec(p, "Background")
        self.pg_sw = {}
        for key, lab in (("bg", "COLOR 1"), ("bg2", "COLOR 2")):
            f = self.irow(s, lab, 8)
            sw = Swatch(f, lambda k=key: self.pick_page_color(k), w=64, h=20)
            sw.pack(side="left")
            self.pg_sw[key] = sw
        self.grad_var = tk.IntVar()
        DarkCheck(s, "Gradient (1 → 2)", self.grad_var, self.toggle_grad, font=("Consolas", 8, "bold")).pack(fill="x", pady=2)
        self.grid_btns(s, [("🖼 Image…", self.set_bgimg), ("✕ Remove", self.clear_bgimg)], 2, px=3)
        self.small_btn(s, "Apply to all pages", self.bg_all, py=5).pack(fill="x", padx=2, pady=(2, 0))
        s = self.sec(p, "Pages")
        self.grid_btns(s, [("＋ New", self.add_page), ("⧉ Duplicate", self.dup_page), ("▲ Up", lambda: self.move_page(-1)),
                           ("▼ Down", lambda: self.move_page(1)), ("🗑 Delete", self.del_page)], 2, px=3)

    def commit_name(self, e=None):
        if self.syncing:
            return
        el = self.cur()
        if el is None:
            return
        v = self.name_var.get().strip()
        if v == el["name"]:
            return
        self.push_undo()
        el["name"] = v
        self._layers_sig = None
        self.sync_layers()

    def _m_jump(self, m):
        for i in range(len(self.pages)):
            m.add_command(label=f"Page {i + 1}", command=lambda i=i: self.set_link(f"#{i + 1}"))

    def _m_sizes(self, m):
        for n in SIZES:
            m.add_command(label=n, command=lambda n=n: self.set_size(*SIZES[n]))
        m.add_separator()
        m.add_command(label="Custom…", command=self.custom_size)

    def bbox_of(self, els):
        x0 = min(min(e["x"], e["x"] + e["w"]) for e in els)
        y0 = min(min(e["y"], e["y"] + e["h"]) for e in els)
        x1 = max(max(e["x"], e["x"] + e["w"]) for e in els)
        y1 = max(max(e["y"], e["y"] + e["h"]) for e in els)
        return x0, y0, x1, y1

    def sync_ui(self):
        self.syncing = True
        try:
            els = self.selels()
            e0 = els[0] if els else None
            el = self.cur()
            pg = self.page
            tx = e0 is not None and e0["t"] in TEXTY
            self.font_btn.config(text=(e0["font"] if tx else "—") + " ▾")
            self.size_ent.set(e0["size"] if tx else None)
            for key, lst in self.swatches.items():
                v = (e0[key] if (e0 and (key != "color" or e0["t"] in TEXTY)) else "")
                for sw in lst:
                    sw.set(v)
            if el is not None:
                for k in "xywh":
                    self.geo[k].set(el[k])
                self.el_name.config(text=f"{TYPE_NAMES[el['t']].upper()}  #{self.sel[0] + 1}")
                self.name_var.set(el["name"])
            elif els:
                x0, y0, x1, y1 = self.bbox_of(els)
                for k, v in zip("xywh", (x0, y0, x1 - x0, y1 - y0)):
                    self.geo[k].set(v)
                self.el_name.config(text=f"{len(els)} ELEMENTS SELECTED")
                self.name_var.set("")
            else:
                for k in "xywh":
                    self.geo[k].set(None)
                self.el_name.config(text="NO SELECTION")
                self.name_var.set("")
            self.op_sl.set((e0["op"] if e0 else 1.0) * 100)
            self.lw_sl.set(e0["lw"] if e0 else 1)
            self.r_sl.set(e0["r"] if e0 else 0)
            self.tx.delete("1.0", "end")
            if el is not None and el["t"] in TEXTY:
                self.tx.insert("1.0", el["text"])
            if e0 is not None:
                self.link_var.set(e0["video"] if e0["t"] == "video" else e0["link"])
                if all(e["t"] in ("image", "video") for e in els):
                    self.replace_btn.pack(fill="x", pady=(4, 0))
                else:
                    self.replace_btn.pack_forget()
            else:
                self.link_var.set("")
                self.replace_btn.pack_forget()
            self.update_link_hint()
            self.pg_sw["bg"].set(pg["bg"])
            self.pg_sw["bg2"].set(pg["bg2"])
            self.grad_var.set(1 if pg["grad"] else 0)
            name = next((n for n, v in SIZES.items() if v == (self.W, self.H)), f"Custom  ({self.W}×{self.H})")
            self.size_btn.config(text=name + " ▾")
        finally:
            self.syncing = False
        self.refresh_toggles()
        self.sync_layers()
        self.update_sel_label()

    def update_sel_label(self):
        els = self.selels()
        self.vp_info.config(text=f"{self.W}×{self.H} pt   ·   page {self.pi + 1}/{len(self.pages)}")
        if not els:
            self.sel_lbl.config(text=f"{len(self.els)} elements on this page")
        else:
            x0, y0, x1, y1 = self.bbox_of(els)
            n = f"{len(els)} elements" if len(els) > 1 else TYPE_NAMES[els[0]["t"]]
            self.sel_lbl.config(text=f"{n}   ·   X {x0:.0f}  Y {y0:.0f}   ·   {x1 - x0:.0f}×{y1 - y0:.0f}")

    def update_link_hint(self):
        e0 = self.selels()[0] if self.selels() else None
        v = self.link_var.get().strip()
        msg, col = "https://…   or   #3  (jump to page 3)", C["text3"]
        if e0 is not None and v.startswith("#"):
            if v[1:].isdigit() and 1 <= int(v[1:]) <= len(self.pages):
                msg, col = f"→ jumps to page {int(v[1:])}", C["ok"]
            else:
                msg, col = "⚠ no such page", C["warn"]
        self.link_hint.config(text=msg, fg=col)

    def relayout_inspector(self):
        return

    def sync_layers(self):
        sig = (tuple((e["t"], e["name"], e["text"][:16], bool(e["link"] or e["video"])) for e in self.els), tuple(self.sel))
        if sig == self._layers_sig:
            return
        self._layers_sig = sig
        inner = self.layers_sf.inner
        for ch in inner.winfo_children():
            ch.destroy()
        if not self.els:
            tk.Label(inner, text="(empty page)", bg=C["inset"], fg=C["text3"], font=("Consolas", 8), pady=6).pack(fill="x")
        for i in range(len(self.els) - 1, -1, -1):
            e = self.els[i]
            on = i in self.sel
            nm = e["name"] or (e["text"].replace("\n", " ")[:18] if e["t"] in TEXTY and e["text"] else TYPE_NAMES[e["t"]])
            lab = f" {TYPE_ICON[e['t']]}  {nm}" + ("   ↗" if (e["link"] or e["video"]) else "")
            row = tk.Label(inner, text=lab, bg=C["accent"] if on else C["inset"], fg="#000000" if on else C["text"],
                           font=("Consolas", 9), anchor="w", pady=3, cursor="hand2")
            row.pack(fill="x")
            row.bind("<Button-1>", lambda ev, i=i: self.layer_click(i, ev))
            row.bind("<Double-Button-1>", lambda ev, i=i: self.rename_layer(i))
            self.layers_sf.bind_wheel(row)

    def layer_click(self, i, ev):
        if ev.state & 0x1:
            self.sel = sorted(set(self.sel) ^ {i})
        else:
            self.sel = [i]
        self.sync_ui()
        self.redraw()

    def rename_layer(self, i):
        e = self.els[i]
        v = ask_string(self.root, "Layer name", "Name:", e["name"] or TYPE_NAMES[e["t"]])
        if v is not None:
            self.push_undo()
            e["name"] = v.strip()
            self.sync_ui()

    def after_change(self, structure=False):
        self.sync_ui()
        self.redraw()
        self.schedule_thumb()

    def set_attr(self, k, v, texty=False, co=None):
        if self.syncing:
            return
        els = [e for e in self.selels() if (not texty or e["t"] in TEXTY)]
        if k == "fill":
            els = [e for e in els if e["t"] in FILLY]
        elif k == "line":
            els = [e for e in els if e["t"] in FILLY + ("line",)]
        if not els or all(e.get(k) == v for e in els):
            return
        if co:
            self.coalesce(co)
        else:
            self.push_undo()
        for e in els:
            e[k] = v
        self.after_change()

    def toggle(self, k):
        els = [e for e in self.selels() if e["t"] in TEXTY]
        if els:
            self.set_attr(k, not els[0][k], texty=True)

    def bump_size(self, d):
        els = [e for e in self.selels() if e["t"] in TEXTY]
        if els:
            self.set_attr("size", float(max(4, min(400, els[0]["size"] + d * (2 if els[0]["size"] >= 24 else 1)))), texty=True, co="size")

    def commit_geom(self, key, v):
        if self.syncing:
            return
        els = self.selels()
        if not els:
            return
        if len(els) == 1:
            el = els[0]
            if key in ("w", "h") and el["t"] != "line":
                v = max(4.0, v)
            if abs(el[key] - v) < 1e-6:
                return
            self.coalesce("geom")
            el[key] = v
        else:
            if key in ("w", "h"):
                return
            x0, y0, _, _ = self.bbox_of(els)
            d = v - (x0 if key == "x" else y0)
            if abs(d) < 1e-6:
                return
            self.coalesce("geom")
            for e in els:
                e[key] += d
        self.after_change()

    def commit_text(self, e=None):
        el = self.cur()
        if self.syncing or el is None or el["t"] not in TEXTY:
            return
        new = self.tx.get("1.0", "end-1c")
        if new == el["text"]:
            return
        self.coalesce("text")
        el["text"] = new
        self.redraw()
        self.schedule_thumb()
        self._layers_sig = None
        self.sync_layers()

    def commit_link(self):
        if self.syncing:
            return
        v = self.link_var.get().strip()
        els = self.selels()
        changed = [e for e in els if e["video" if e["t"] == "video" else "link"] != v]
        if not changed:
            self.update_link_hint()
            return
        self.push_undo()
        for e in changed:
            e["video" if e["t"] == "video" else "link"] = v
        self.after_change()

    def set_link(self, v):
        self.link_var.set(v)
        self.commit_link()

    def clear_link(self):
        self.set_link("")

    def pick_attr_color(self, key):
        e0 = self.selels()[0] if self.selels() else None
        if e0 is None:
            self.status("Select an element first.", "warn")
            return
        c = pick_color(self.root, e0[key] or "#ffffff", {"color": "Text color", "fill": "Fill", "line": "Border"}[key])
        if c:
            self.set_color(key, c)

    def set_color(self, key, c):
        self.set_attr(key, c, texty=(key == "color"))

    def pick_page_color(self, key):
        c = pick_color(self.root, self.page[key], "Background color")
        if c:
            self.push_undo()
            self.page[key] = c
            self.after_change()

    def toggle_grad(self):
        if self.syncing:
            return
        self.push_undo()
        self.page["grad"] = bool(self.grad_var.get())
        self.after_change()

    def set_size(self, w, h):
        if (w, h) == (self.W, self.H):
            return
        self.push_undo()
        self.W, self.H = w, h
        self.auto_fit = True
        self.refresh_thumbs()
        self.after_change()

    def custom_size(self):
        v = ask_string(self.root, "Custom size", "Width×Height in points (e.g. 1280x720):", f"{self.W}x{self.H}")
        if not v:
            return
        m = re.match(r"\s*(\d+)\s*[x×*,; ]\s*(\d+)\s*$", v)
        if not m or not (50 <= int(m.group(1)) <= 5000 and 50 <= int(m.group(2)) <= 5000):
            self.status("Invalid size (50–5000).", "err")
            return
        self.set_size(int(m.group(1)), int(m.group(2)))

    def bg_all(self):
        self.push_undo()
        for p in self.pages:
            for k in ("bg", "bg2", "grad", "bgimg"):
                p[k] = self.page[k]
        self.status("Background copied to all pages.")
        self.refresh_thumbs()
        self.redraw()

    def fit_zoom(self, cw, ch):
        return max(0.05, min(8.0, min((cw - 90) / self.W, (ch - 90) / self.H)))

    def zoom_at(self, f, mx=None, my=None):
        cw, ch = self.cv.winfo_width(), self.cv.winfo_height()
        if mx is None:
            mx, my = cw / 2, ch / 2
        nz = max(0.05, min(8.0, self.z * f))
        px, py = (mx - self.ox) / self.z, (my - self.oy) / self.z
        self.auto_fit = False
        self.zoom = nz
        self.panx = mx - px * nz - (cw - self.W * nz) / 2
        self.pany = my - py * nz - (ch - self.H * nz) / 2
        self.redraw()

    def zoom_by(self, f):
        self.zoom_at(f)

    def set_zoom(self, v):
        self.zoom_at(v / self.z)

    def fit(self):
        self.auto_fit = True
        self.panx = self.pany = 0.0
        self.redraw()

    def on_wheel(self, e):
        up = e.num == 4 or getattr(e, "delta", 0) > 0
        self.zoom_at(1.12 if up else 1 / 1.12, e.x, e.y)

    def on_pan_start(self, e):
        self.cv.focus_set()
        self.drag = dict(mode="pan", sx=e.x, sy=e.y, p=(self.panx, self.pany), moved=True)

    def on_pan(self, e):
        d = self.drag
        if d and d["mode"] == "pan":
            self.auto_fit = False
            self.panx = d["p"][0] + e.x - d["sx"]
            self.pany = d["p"][1] + e.y - d["sy"]
            self.redraw()

    def add_asset(self, data):
        aid = uuid.uuid4().hex[:12]
        self.assets[aid] = base64.b64encode(data).decode()
        return aid

    def pil(self, aid, op=1.0):
        if aid not in self.assets:
            return None
        if aid not in self.pil_cache:
            try:
                self.pil_cache[aid] = Image.open(io.BytesIO(base64.b64decode(self.assets[aid]))).convert("RGBA")
            except Exception:
                return None
        im = self.pil_cache[aid]
        if op < 0.999:
            k = (aid, round(op, 2))
            if k not in self.pil_cache:
                if len(self.pil_cache) > 120:
                    for kk in [kk for kk in self.pil_cache if isinstance(kk, tuple)]:
                        del self.pil_cache[kk]
                a = im.getchannel("A").point(lambda v: int(v * op))
                im2 = im.copy()
                im2.putalpha(a)
                self.pil_cache[k] = im2
            return self.pil_cache[k]
        return im

    def ask_image(self, title="Choose image"):
        p = filedialog.askopenfilename(parent=self.root, title=title, initialdir=self.last_dir, filetypes=[
            ("Images", "*.png *.jpg *.jpeg *.bmp *.gif *.webp *.tif *.tiff"), ("All files", "*.*")])
        if not p:
            return None
        self.last_dir = os.path.dirname(p)
        try:
            with open(p, "rb") as f:
                data = f.read()
            Image.open(io.BytesIO(data)).verify()
        except Exception as ex:
            self.status(f"Image error: {ex}", "err")
            return None
        return data

    def add_image(self):
        data = self.ask_image()
        if not data:
            return
        aid = self.add_asset(data)
        im = self.pil(aid)
        k = min(self.W * 0.6 / im.width, self.H * 0.6 / im.height, 1.0 if max(im.size) < 2000 else 0.5)
        w, h = im.width * k, im.height * k
        self.push_undo()
        self.els.append(new_el("image", self.fams[0], asset=aid, w=w, h=h, x=(self.W - w) / 2, y=(self.H - h) / 2))
        self.sel = [len(self.els) - 1]
        self.after_change()

    def add_video(self):
        ch = _run_dialog(self.root, CustomMessageBox(self.root, "Video", "Video file on this computer, or a web address (e.g. YouTube)?\n\n"
                                                     "In the PDF the video becomes a poster image + clickable link.",
                                                     "triple", yes_text="File", no_text="Web URL", cancel_text="Cancel"))
        if ch is None:
            return
        link, poster = "", None
        if ch:
            p = filedialog.askopenfilename(parent=self.root, title="Video", initialdir=self.last_dir,
                                           filetypes=[("Video", "*.mp4 *.mov *.avi *.mkv *.webm *.wmv"), ("All files", "*.*")])
            if not p:
                return
            self.last_dir = os.path.dirname(p)
            link = Path(p).resolve().as_uri()
            if cv2 is not None:
                try:
                    cap = cv2.VideoCapture(p)
                    cap.set(cv2.CAP_PROP_POS_MSEC, 1000)
                    ok, fr = cap.read()
                    if not ok:
                        cap.set(cv2.CAP_PROP_POS_FRAMES, 0)
                        ok, fr = cap.read()
                    cap.release()
                    if ok:
                        im = Image.fromarray(cv2.cvtColor(fr, cv2.COLOR_BGR2RGB))
                        buf = io.BytesIO()
                        im.save(buf, "PNG")
                        poster = self.add_asset(buf.getvalue())
                except Exception:
                    poster = None
        else:
            link = ask_string(self.root, "Video URL", "Video web address:", "https://")
            if not link or link.strip() in ("", "https://"):
                return
            link = link.strip()
            if "://" not in link:
                link = "https://" + link
        self.push_undo()
        el = new_el("video", self.fams[0], video=link, asset=poster or "", x=(self.W - 480) / 2, y=(self.H - 270) / 2)
        self.els.append(el)
        self.sel = [len(self.els) - 1]
        self.after_change()
        extra = "" if (poster or not ch) else "  (without opencv-python there is no automatic poster - use \"Replace image / poster\" to set one)"
        self.status("Video: a poster image + clickable link goes into the PDF." + extra)

    def replace_asset(self):
        els = [e for e in self.selels() if e["t"] in ("image", "video")]
        if not els:
            self.status("Select an image or a video.", "warn")
            return
        data = self.ask_image()
        if data:
            aid = self.add_asset(data)
            self.push_undo()
            for e in els:
                e["asset"] = aid
            self.after_change()

    def set_bgimg(self):
        data = self.ask_image("Background image")
        if data:
            self.push_undo()
            self.page["bgimg"] = self.add_asset(data)
            self.after_change()

    def clear_bgimg(self):
        if self.page["bgimg"]:
            self.push_undo()
            self.page["bgimg"] = ""
            self.after_change()

    def add_el(self, kind):
        self.push_undo()
        el = new_el(kind, self.fams[0])
        el["x"], el["y"] = (self.W - el["w"]) / 2, (self.H - el["h"]) / 2
        self.els.append(el)
        self.sel = [len(self.els) - 1]
        self.after_change()
        if kind == "text":
            self.root.after(30, lambda: self.edit_text(len(self.els) - 1) if self.els and self.els[-1] is el else None)

    def del_el(self):
        if not self.sel:
            return
        self.push_undo()
        for i in sorted(self.sel, reverse=True):
            if 0 <= i < len(self.els):
                del self.els[i]
        self.sel = []
        self.after_change()

    def dup_el(self):
        els = self.selels()
        if not els:
            return
        self.push_undo()
        new = []
        for e in els:
            c = copy.deepcopy(e)
            c["x"] += 20
            c["y"] += 20
            self.els.append(c)
            new.append(len(self.els) - 1)
        self.sel = new
        self.after_change()

    def select_all(self):
        self.sel = list(range(len(self.els)))
        self.sync_ui()
        self.redraw()

    def copy(self):
        els = self.selels()
        if not els:
            return
        aids = {e["asset"] for e in els if e["asset"]}
        self.clip = dict(els=copy.deepcopy(els), assets={a: self.assets[a] for a in aids if a in self.assets}, n=0)
        self.status(f"{len(els)} element(s) copied.")

    def cut(self):
        self.copy()
        self.del_el()

    def paste(self):
        if not self.clip:
            return
        self.clip["n"] += 1
        off = 20 * self.clip["n"]
        self.push_undo()
        for a, b in self.clip["assets"].items():
            self.assets.setdefault(a, b)
        new = []
        for e in copy.deepcopy(self.clip["els"]):
            e["x"] += off
            e["y"] += off
            self.els.append(e)
            new.append(len(self.els) - 1)
        self.sel = new
        self.after_change()

    def z_move(self, how):
        sel = sorted(self.sel)
        if not sel:
            return
        els = self.els
        chosen = [els[i] for i in sel]
        if how in ("front", "back"):
            others = [e for i, e in enumerate(els) if i not in sel]
            new = others + chosen if how == "front" else chosen + others
        else:
            new = list(els)
            S = set(sel)
            order = sorted(sel, reverse=True) if how == 1 else sorted(sel)
            for i in order:
                j = i + how
                if 0 <= j < len(new) and j not in S:
                    new[i], new[j] = new[j], new[i]
                    S.discard(i)
                    S.add(j)
        if all(a is b for a, b in zip(new, els)):
            return
        self.push_undo()
        self.page["els"] = new
        self.sel = [k for k, e in enumerate(new) if any(e is c for c in chosen)]
        self.after_change()

    def align(self, how):
        els = self.selels()
        if not els:
            return
        if len(els) == 1:
            bx0, by0, bx1, by1 = 0, 0, self.W, self.H
        else:
            bx0, by0, bx1, by1 = self.bbox_of(els)
        self.push_undo()
        for e in els:
            ex0, ey0, ex1, ey1 = self.bbox_of([e])
            dx = dy = 0
            if how == "l":
                dx = bx0 - ex0
            elif how == "c":
                dx = (bx0 + bx1) / 2 - (ex0 + ex1) / 2
            elif how == "r":
                dx = bx1 - ex1
            elif how == "t":
                dy = by0 - ey0
            elif how == "m":
                dy = (by0 + by1) / 2 - (ey0 + ey1) / 2
            elif how == "b":
                dy = by1 - ey1
            e["x"] += dx
            e["y"] += dy
        self.after_change()

    def distribute(self, axis):
        els = self.selels()
        if len(els) < 3:
            self.status("Distribute needs at least 3 elements.", "warn")
            return
        i0, i1 = (0, 2) if axis == "h" else (1, 3)
        info = [(e, self.bbox_of([e])) for e in els]
        info.sort(key=lambda t: (t[1][i0] + t[1][i1]) / 2)
        start, end = info[0][1][i0], info[-1][1][i1]
        total = sum(b[i1] - b[i0] for _, b in info)
        gap = (end - start - total) / (len(info) - 1)
        self.push_undo()
        pos = start
        for e, b in info:
            d = pos - b[i0]
            e["x" if axis == "h" else "y"] += d
            pos += (b[i1] - b[i0]) + gap
        self.after_change()

    def goto(self, i):
        if not (0 <= i < len(self.pages)):
            return
        self.end_edit()
        self.pi = i
        self.sel = []
        self.mark_thumbs()
        self.sync_ui()
        self.redraw()

    def add_page(self):
        self.push_undo()
        self.pages.insert(self.pi + 1, self.blank_page())
        self.pi += 1
        self.sel = []
        self.refresh_thumbs()
        self.after_change()

    def dup_page(self):
        self.push_undo()
        self.pages.insert(self.pi + 1, copy.deepcopy(self.page))
        self.pi += 1
        self.sel = []
        self.refresh_thumbs()
        self.after_change()

    def del_page(self):
        if len(self.pages) <= 1:
            self.status("The last page cannot be deleted.", "warn")
            return
        self.push_undo()
        del self.pages[self.pi]
        self.pi = max(0, self.pi - 1)
        self.sel = []
        self.refresh_thumbs()
        self.after_change()

    def move_page(self, d):
        self.move_page_to(self.pi, self.pi + d)

    def move_page_to(self, i, j):
        if not (0 <= j < len(self.pages)) or i == j:
            return
        self.push_undo()
        pg = self.pages.pop(i)
        self.pages.insert(j, pg)
        self.pi = j
        self.refresh_thumbs()
        self.sync_ui()
        self.redraw()

    def refresh_thumbs(self):
        self.cache_thumb.trim()
        inner = self.th.inner
        for ch in inner.winfo_children():
            ch.destroy()
        self.thumbs = []
        tw = 148
        th = max(24, int(tw * self.H / self.W))
        for i, pg in enumerate(self.pages):
            f = tk.Frame(inner, bg=C["panel2"])
            f.pack(fill="x", pady=(6, 0), padx=4)
            num = tk.Label(f, text=str(i + 1), bg=C["panel2"], fg=C["text2"], font=("Consolas", 8, "bold"), width=2, anchor="ne")
            num.pack(side="left", anchor="n", pady=2)
            c = tk.Canvas(f, width=tw, height=th, bg="#ffffff", highlightthickness=2, highlightbackground=C["border"],
                          cursor="hand2")
            c.pack(side="left")
            for w in (c, num, f):
                w.bind("<ButtonPress-1>", lambda e, i=i: self.goto(i))
                w.bind("<ButtonRelease-1>", lambda e, i=i: self.thumb_release(i, e))
                w.bind("<ButtonPress-3>", lambda e, i=i: self.page_menu(i, e))
                self.th.bind_wheel(w)
            self.thumbs.append((f, num, c))
        for i in range(len(self.pages)):
            self.paint_thumb(i)
        self.mark_thumbs()

    def paint_thumb(self, i):
        if not (0 <= i < len(self.thumbs)):
            return
        c = self.thumbs[i][2]
        c.delete("all")
        tw = int(c["width"])
        self.paint_page(c, self.pages[i], tw / self.W, 0, 0, self.cache_thumb, badges=False)

    def schedule_thumb(self):
        if self._thumb_job is None:
            def go():
                self._thumb_job = None
                self.paint_thumb(self.pi)
            self._thumb_job = self.root.after(140, go)

    def mark_thumbs(self):
        for i, (f, num, c) in enumerate(self.thumbs):
            cur = i == self.pi
            c.config(highlightbackground=C["accent"] if cur else C["border"])
            num.config(fg=C["accent"] if cur else C["text2"])
        if 0 <= self.pi < len(self.thumbs):
            self.root.after(20, self.see_thumb)

    def see_thumb(self):
        try:
            f = self.thumbs[self.pi][0]
            cv = self.th.cv
            total = max(1, self.th.inner.winfo_height())
            y0, y1 = f.winfo_y(), f.winfo_y() + f.winfo_height()
            top, bot = cv.yview()
            if y0 < top * total:
                cv.yview_moveto(max(0, (y0 - 6) / total))
            elif y1 > bot * total:
                cv.yview_moveto(max(0, (y1 + 6) / total - (bot - top)))
        except Exception:
            pass

    def thumb_release(self, i, e):
        w = self.root.winfo_containing(e.x_root, e.y_root)
        for j, (f, num, c) in enumerate(self.thumbs):
            if w is c or w is f or w is num:
                if j != i:
                    self.move_page_to(i, j)
                return

    def page_menu(self, i, e):
        self.goto(i)
        m = CtxMenu(self)
        m.add_command(label="New page after", command=self.add_page)
        m.add_command(label="Duplicate", command=self.dup_page)
        m.add_separator()
        m.add_command(label="Move up", command=lambda: self.move_page(-1))
        m.add_command(label="Move down", command=lambda: self.move_page(1))
        m.add_separator()
        m.add_command(label="Delete", command=self.del_page)
        m.popup(e.x_root, e.y_root)

    def tkfont(self, el, px):
        key = (el["font"], el["b"], el["i"], px)
        if key not in self.fonts:
            if len(self.fonts) > 2000:
                self.fonts.clear()
            self.fonts[key] = tkfont.Font(family=el["font"], size=-max(1, px), weight="bold" if el["b"] else "normal",
                                          slant="italic" if el["i"] else "roman")
        return self.fonts[key]

    def img_for(self, cache, aid, w, h, op):
        im = self.pil(aid, op)
        if im is None:
            return None
        return cache.get((aid, round(op, 2)), im, w, h)

    def paint_page(self, cv, pg, z, ox, oy, cache, badges=False, hide=None):
        W, H = self.W, self.H
        done = False
        if pg["bgimg"] and self.pil(pg["bgimg"]):
            ph = cache.get((pg["bgimg"], 1.0), self.pil(pg["bgimg"]), W * z, H * z)
            if ph:
                cv.create_image(ox, oy, anchor="nw", image=ph)
                done = True
        if not done and pg["grad"]:
            n = 48 if z > 0.4 else 14
            for i in range(n):
                cv.create_rectangle(ox, oy + H * z * i / n, ox + W * z, oy + H * z * (i + 1) / n + 1,
                                    fill=lerp_hex(pg["bg"], pg["bg2"], i / (n - 1)), outline="")
            done = True
        if not done:
            cv.create_rectangle(ox, oy, ox + W * z, oy + H * z, fill=pg["bg"], outline="")
        for el in pg["els"]:
            self.paint_el(cv, el, z, ox, oy, pg["bg"], cache, badges, hide is el)

    def paint_el(self, cv, el, z, ox, oy, bg, cache, badges=False, hide_text=False):
        x, y, w, h = ox + el["x"] * z, oy + el["y"] * z, el["w"] * z, el["h"] * z
        t, op = el["t"], el["op"]
        fill = blend(el["fill"], bg, op) if el["fill"] else ""
        line = blend(el["line"], bg, op) if el["line"] else ""
        lw = max(1.0, el["lw"] * z) if el["line"] else 0
        if t == "line":
            cv.create_line(x, y, x + w, y + h, fill=line or "#888888", width=max(1, el["lw"] * z), capstyle="round")
            return
        if t == "image":
            ph = self.img_for(cache, el["asset"], abs(w), abs(h), op)
            if ph:
                cv.create_image(x, y, anchor="nw", image=ph)
            else:
                cv.create_rectangle(x, y, x + w, y + h, fill="#333333", outline="#666666")
                cv.create_text(x + w / 2, y + h / 2, text="▨", fill="#777777", font=("Consolas", max(8, int(min(w, h) * 0.3))))
            return
        if t == "video":
            ph = self.img_for(cache, el["asset"], abs(w), abs(h), op) if el["asset"] else None
            if ph:
                cv.create_image(x, y, anchor="nw", image=ph)
            else:
                cv.create_rectangle(x, y, x + w, y + h, fill="#101010", outline="#555555")
            cx, cy, r = x + w / 2, y + h / 2, min(w, h) * 0.14
            cv.create_oval(cx - r, cy - r, cx + r, cy + r, fill="#ffffff", outline="")
            cv.create_polygon(cx - r * 0.35, cy - r * 0.55, cx - r * 0.35, cy + r * 0.55, cx + r * 0.6, cy, fill="#222222", outline="")
            return
        if t == "ellipse":
            cv.create_oval(x, y, x + w, y + h, fill=fill, outline=line, width=lw)
        else:
            rr = min(el["r"] * z, w / 2, h / 2)
            if (t in ("rect", "button", "text")) and (fill or line):
                if rr > 1:
                    self.round_rect(cv, x, y, x + w, y + h, rr, fill, line, lw)
                else:
                    cv.create_rectangle(x, y, x + w, y + h, fill=fill, outline=line, width=lw)
        if el["text"] and not hide_text:
            lines, lh, bt = text_layout(el)
            size, pad = el["size"], 6
            f = self.tkfont(el, max(1, int(round(size * z))))
            asc = f.metrics("ascent")
            col = blend(el["color"], bg, op)
            for k, ln in enumerate(lines):
                if not ln:
                    continue
                ty = y + (bt + k * lh + 0.88 * size) * z - asc
                if el["align"] == "center":
                    cv.create_text(x + w / 2, ty, anchor="n", text=ln, fill=col, font=f)
                elif el["align"] == "right":
                    cv.create_text(x + w - pad * z, ty, anchor="ne", text=ln, fill=col, font=f)
                else:
                    cv.create_text(x + pad * z, ty, anchor="nw", text=ln, fill=col, font=f)
        if badges and (el["link"] and t != "video"):
            cv.create_text(x + w - 3, y + 2, anchor="ne", text="↗", fill=C["accent"], font=("Consolas", max(8, int(10 * z ** 0.3)), "bold"))

    def round_rect(self, cv, x0, y0, x1, y1, r, fill, outline, width):
        pts = [x0 + r, y0, x1 - r, y0, x1, y0, x1, y0 + r, x1, y1 - r, x1, y1, x1 - r, y1, x0 + r, y1, x0, y1, x0, y1 - r,
               x0, y0 + r, x0, y0]
        cv.create_polygon(pts, smooth=True, fill=fill, outline=outline, width=width)

    def redraw(self):
        cv = self.cv
        cw, ch = cv.winfo_width(), cv.winfo_height()
        if cw < 60 or ch < 60:
            return
        if self.auto_fit:
            self.zoom = self.fit_zoom(cw, ch)
            self.panx = self.pany = 0.0
        z = self.z = self.zoom
        self.ox = (cw - self.W * z) / 2 + self.panx
        self.oy = (ch - self.H * z) / 2 + self.pany
        ox, oy = self.ox, self.oy
        cv.delete("all")
        self.cache_main.trim()
        cv.create_rectangle(ox + 6, oy + 6, ox + self.W * z + 6, oy + self.H * z + 6, fill=blend("#000000", C["view"], 0.45), outline="")
        self.paint_page(cv, self.page, z, ox, oy, self.cache_main, badges=True, hide=self.editing)
        self.mask_outside(cv, ox, oy, z, cw, ch, C["view"], "gray75")
        if self.grid_on and GRID * z >= 6:
            gx = 0
            while gx <= self.W:
                cv.create_line(ox + gx * z, oy, ox + gx * z, oy + self.H * z, fill="#8a96a8", dash=(1, 4))
                gx += GRID
            gy = 0
            while gy <= self.H:
                cv.create_line(ox, oy + gy * z, ox + self.W * z, oy + gy * z, fill="#8a96a8", dash=(1, 4))
                gy += GRID
        cv.create_rectangle(ox, oy, ox + self.W * z, oy + self.H * z, outline=C["border2"])
        self.draw_sel()
        for kind, v in self.guides:
            if kind == "v":
                cv.create_line(ox + v * z, oy - 10, ox + v * z, oy + self.H * z + 10, fill=C["guide"], dash=(4, 3))
            else:
                cv.create_line(ox - 10, oy + v * z, ox + self.W * z + 10, oy + v * z, fill=C["guide"], dash=(4, 3))
        if self.marquee:
            x0, y0, x1, y1 = self.marquee
            cv.create_rectangle(x0, y0, x1, y1, outline=C["accent"], dash=(3, 3), fill="", width=1)
        if self.editor is not None and self.editing is not None:
            el = self.editing
            cv.create_window(ox + el["x"] * z, oy + el["y"] * z, anchor="nw", window=self.editor,
                             width=max(30, int(el["w"] * z)), height=max(24, int(el["h"] * z)))
        self.zl.config(text=f"{int(round(z * 100))}%")

    def mask_outside(self, cv, ox, oy, z, cw, ch, color, stipple=None):
        x1, y1 = ox + self.W * z, oy + self.H * z
        kw = dict(fill=color, outline="")
        if stipple:
            kw["stipple"] = stipple
        for r in ((0, 0, cw, oy), (0, y1, cw, ch), (0, oy, ox, y1), (x1, oy, cw, y1)):
            if r[2] > r[0] and r[3] > r[1]:
                cv.create_rectangle(*r, **kw)

    def handles(self, el):
        z = self.z
        x, y, w, h = self.ox + el["x"] * z, self.oy + el["y"] * z, el["w"] * z, el["h"] * z
        if el["t"] == "line":
            return {"p0": (x, y), "p1": (x + w, y + h)}
        xs = {"w": x, "c": x + w / 2, "e": x + w}
        ys = {"n": y, "c": y + h / 2, "s": y + h}
        out = {}
        for a in "nsc":
            for b in "wec":
                if a == "c" and b == "c":
                    continue
                out[a + b] = (xs[b], ys[a])
        return out

    def draw_sel(self):
        els = self.selels()
        if not els:
            return
        z, cv = self.z, self.cv
        for e in els:
            if e is self.editing:
                continue
            x0, y0, x1, y1 = self.bbox_of([e])
            if e["t"] != "line":
                cv.create_rectangle(self.ox + x0 * z, self.oy + y0 * z, self.ox + x1 * z, self.oy + y1 * z,
                                    outline=C["accent"], dash=(5, 3) if len(els) == 1 else (), width=1)
        if len(els) > 1:
            x0, y0, x1, y1 = self.bbox_of(els)
            cv.create_rectangle(self.ox + x0 * z - 3, self.oy + y0 * z - 3, self.ox + x1 * z + 3, self.oy + y1 * z + 3,
                                outline=C["accent2"], dash=(6, 3))
        elif self.editing is None:
            for k, (hx, hy) in self.handles(els[0]).items():
                cv.create_rectangle(hx - 4, hy - 4, hx + 4, hy + 4, fill=C["accent"], outline="#000000")

    def to_page(self, mx, my):
        return (mx - self.ox) / self.z, (my - self.oy) / self.z

    def hit_handle(self, mx, my):
        el = self.cur()
        if el is None or self.editing is not None:
            return None
        for k, (hx, hy) in self.handles(el).items():
            if abs(mx - hx) <= 7 and abs(my - hy) <= 7:
                return k
        return None

    @staticmethod
    def seg_dist(px, py, x0, y0, x1, y1):
        dx, dy = x1 - x0, y1 - y0
        L = dx * dx + dy * dy
        t = 0 if L == 0 else max(0, min(1, ((px - x0) * dx + (py - y0) * dy) / L))
        return math.hypot(px - (x0 + t * dx), py - (y0 + t * dy))

    def hit_el(self, mx, my):
        px, py = self.to_page(mx, my)
        for i in range(len(self.els) - 1, -1, -1):
            e = self.els[i]
            if e["t"] == "line":
                if self.seg_dist(px, py, e["x"], e["y"], e["x"] + e["w"], e["y"] + e["h"]) <= max(6 / self.z, e["lw"] / 2 + 2 / self.z):
                    return i
                continue
            x0, y0, x1, y1 = self.bbox_of([e])
            if x0 <= px <= x1 and y0 <= py <= y1:
                return i
        return None

    def on_press(self, e):
        self.cv.focus_set()
        self.end_edit()
        CtxMenu.close_all()
        if self.space:
            self.drag = dict(mode="pan", sx=e.x, sy=e.y, p=(self.panx, self.pany), moved=True)
            return
        shift = bool(e.state & 0x1)
        hk = self.hit_handle(e.x, e.y)
        if hk:
            el = self.cur()
            self.drag = dict(mode="resize", h=hk, sx=e.x, sy=e.y, o=(el["x"], el["y"], el["w"], el["h"]), moved=False)
            return
        i = self.hit_el(e.x, e.y)
        if i is None:
            if not shift and self.sel:
                self.sel = []
                self.sync_ui()
                self.redraw()
            self.drag = dict(mode="marquee", sx=e.x, sy=e.y, base=list(self.sel) if shift else [], moved=False)
            return
        if shift:
            if i in self.sel:
                self.sel.remove(i)
                self.sync_ui()
                self.redraw()
                self.drag = None
                return
            self.sel = sorted(self.sel + [i])
        elif i not in self.sel:
            self.sel = [i]
        self.sync_ui()
        self.redraw()
        self.drag = dict(mode="move", sx=e.x, sy=e.y, orig=[(j, self.els[j]["x"], self.els[j]["y"]) for j in self.sel],
                         bb=self.bbox_of(self.selels()), moved=False)

    def snap_move(self, dx, dy, bb):
        x0, y0, x1, y1 = bb
        w, h = x1 - x0, y1 - y0
        nx0, ny0 = x0 + dx, y0 + dy
        self.guides = []
        if not self.snap_on:
            return dx, dy
        thr = 7 / self.z
        cx, cy = [0, self.W / 2, self.W], [0, self.H / 2, self.H]
        sel = set(self.sel)
        for i, e in enumerate(self.els):
            if i in sel:
                continue
            a0, b0, a1, b1 = self.bbox_of([e])
            cx += [a0, (a0 + a1) / 2, a1]
            cy += [b0, (b0 + b1) / 2, b1]
        bx = by = None
        for edge in (nx0, nx0 + w / 2, nx0 + w):
            for c in cx:
                d = c - edge
                if abs(d) < thr and (bx is None or abs(d) < abs(bx[0])):
                    bx = (d, c)
        for edge in (ny0, ny0 + h / 2, ny0 + h):
            for c in cy:
                d = c - edge
                if abs(d) < thr and (by is None or abs(d) < abs(by[0])):
                    by = (d, c)
        if bx:
            dx += bx[0]
            self.guides.append(("v", bx[1]))
        elif self.grid_on:
            dx += round(nx0 / GRID) * GRID - nx0
        if by:
            dy += by[0]
            self.guides.append(("h", by[1]))
        elif self.grid_on:
            dy += round(ny0 / GRID) * GRID - ny0
        return dx, dy

    def on_drag(self, e):
        d = self.drag
        if not d:
            return
        if d["mode"] == "pan":
            self.on_pan(e)
            return
        if d["mode"] == "marquee":
            if not d["moved"] and abs(e.x - d["sx"]) + abs(e.y - d["sy"]) < 4:
                return
            d["moved"] = True
            self.marquee = (d["sx"], d["sy"], e.x, e.y)
            self.redraw()
            return
        if not d["moved"]:
            if abs(e.x - d["sx"]) + abs(e.y - d["sy"]) < 3:
                return
            self.push_undo()
            d["moved"] = True
        dx, dy = (e.x - d["sx"]) / self.z, (e.y - d["sy"]) / self.z
        shift = bool(e.state & 0x1)
        if d["mode"] == "move":
            if shift:
                if abs(dx) > abs(dy):
                    dy = 0
                else:
                    dx = 0
            dx, dy = self.snap_move(dx, dy, d["bb"])
            for j, ox_, oy_ in d["orig"]:
                if j < len(self.els):
                    self.els[j]["x"], self.els[j]["y"] = ox_ + dx, oy_ + dy
        else:
            el = self.cur()
            if el is None:
                return
            x, y, w, h = d["o"]
            hk = d["h"]
            if hk == "p0":
                el["x"], el["y"], el["w"], el["h"] = x + dx, y + dy, w - dx, h - dy
                if shift:
                    ang = round(math.atan2(el["h"], el["w"]) / (math.pi / 4)) * (math.pi / 4)
                    L = math.hypot(el["w"], el["h"])
                    el["x"], el["y"] = x + w - L * math.cos(ang), y + h - L * math.sin(ang)
                    el["w"], el["h"] = L * math.cos(ang), L * math.sin(ang)
            elif hk == "p1":
                el["w"], el["h"] = w + dx, h + dy
                if shift:
                    ang = round(math.atan2(el["h"], el["w"]) / (math.pi / 4)) * (math.pi / 4)
                    L = math.hypot(el["w"], el["h"])
                    el["w"], el["h"] = L * math.cos(ang), L * math.sin(ang)
            else:
                if "e" in hk:
                    w = w + dx
                if "w" in hk:
                    x, w = x + dx, w - dx
                if "s" in hk:
                    h = h + dy
                if "n" in hk:
                    y, h = y + dy, h - dy
                corner = hk[0] != "c" and hk[1] != "c"
                if corner and ((el["t"] == "image") != shift) and d["o"][2] > 0 and d["o"][3] > 0:
                    sw_, sh_ = w / d["o"][2], h / d["o"][3]
                    s = sw_ if abs(sw_ - 1) > abs(sh_ - 1) else sh_
                    w, h = d["o"][2] * s, d["o"][3] * s
                    if "w" in hk:
                        x = d["o"][0] + d["o"][2] - w
                    if "n" in hk:
                        y = d["o"][1] + d["o"][3] - h
                if w < 4:
                    if "w" in hk:
                        x -= 4 - w
                    w = 4
                if h < 4:
                    if "n" in hk:
                        y -= 4 - h
                    h = 4
                el["x"], el["y"], el["w"], el["h"] = x, y, w, h
        self.update_sel_label()
        self.redraw()

    def on_release(self, e):
        d = self.drag
        self.drag = None
        self.guides = []
        if d and d["mode"] == "marquee":
            if d["moved"] and self.marquee:
                x0, y0, x1, y1 = self.marquee
                ax, ay = self.to_page(min(x0, x1), min(y0, y1))
                bx, by = self.to_page(max(x0, x1), max(y0, y1))
                hit = []
                for i, el in enumerate(self.els):
                    a0, b0, a1, b1 = self.bbox_of([el])
                    if a1 >= ax and a0 <= bx and b1 >= ay and b0 <= by:
                        hit.append(i)
                self.sel = sorted(set(d["base"]) | set(hit))
            self.marquee = None
            self.sync_ui()
            self.redraw()
            return
        if d and d["mode"] in ("move", "resize") and d["moved"]:
            self.sync_ui()
            self.schedule_thumb()
        self.redraw()

    def on_motion(self, e):
        if self.space:
            self.cv.config(cursor="fleur")
            return
        c = "arrow"
        hk = self.hit_handle(e.x, e.y)
        if hk:
            c = {"nw": "top_left_corner", "se": "bottom_right_corner", "ne": "top_right_corner", "sw": "bottom_left_corner",
                 "nc": "sb_v_double_arrow", "sc": "sb_v_double_arrow", "cw": "sb_h_double_arrow",
                 "ce": "sb_h_double_arrow"}.get(hk, "crosshair")
        elif self.hit_el(e.x, e.y) is not None:
            c = "fleur"
        self.cv.config(cursor=c)

    def on_dbl(self, e):
        i = self.hit_el(e.x, e.y)
        if i is not None and self.els[i]["t"] in TEXTY:
            self.edit_text(i)

    def on_rclick(self, e):
        self.cv.focus_set()
        self.end_edit()
        i = self.hit_el(e.x, e.y)
        if i is not None and i not in self.sel:
            self.sel = [i]
            self.sync_ui()
            self.redraw()
        elif i is None and self.sel:
            self.sel = []
            self.sync_ui()
            self.redraw()
        m = CtxMenu(self)
        if self.sel:
            if any(x["t"] in TEXTY for x in self.selels()) and len(self.sel) == 1:
                m.add_command(label="Edit text              F2", command=self.edit_selected)
                m.add_separator()
            m.add_command(label="Cut                    Ctrl+X", command=self.cut)
            m.add_command(label="Copy                   Ctrl+C", command=self.copy)
            m.add_command(label="Paste                  Ctrl+V", command=self.paste)
            m.add_command(label="Duplicate              Ctrl+D", command=self.dup_el)
            m.add_command(label="Delete                 Del", command=self.del_el)
            m.add_separator()
            zm = CtxMenu(self)
            zm.add_command(label="Bring to front", command=lambda: self.z_move("front"))
            zm.add_command(label="Bring forward", command=lambda: self.z_move(1))
            zm.add_command(label="Send backward", command=lambda: self.z_move(-1))
            zm.add_command(label="Send to back", command=lambda: self.z_move("back"))
            m.add_cascade(label="Order", menu=zm)
            am = CtxMenu(self)
            for lab, how in (("Left", "l"), ("Center horizontally", "c"), ("Right", "r"), ("Top", "t"),
                             ("Center vertically", "m"), ("Bottom", "b")):
                am.add_command(label=lab, command=lambda h=how: self.align(h))
            m.add_cascade(label="Align", menu=am)
            m.add_separator()
            m.add_command(label="Set link…", command=self.ask_link)
        else:
            m.add_command(label="Paste                  Ctrl+V", command=self.paste)
            m.add_command(label="Select all             Ctrl+A", command=self.select_all)
            m.add_separator()
            im = CtxMenu(self)
            for k in ("text", "rect", "ellipse", "line", "button"):
                im.add_command(label=f"{TYPE_ICON[k]}  {TYPE_NAMES[k]}", command=lambda k=k: self.add_el(k))
            im.add_command(label="🖼  Image…", command=self.add_image)
            im.add_command(label="🎬  Video…", command=self.add_video)
            m.add_cascade(label="Insert", menu=im)
            m.add_command(label="New page               Ctrl+M", command=self.add_page)
            m.add_separator()
            m.add_command(label="Fit to window          Ctrl+0", command=self.fit)
        m.popup(e.x_root, e.y_root)

    def ask_link(self):
        e0 = self.selels()[0] if self.selels() else None
        if e0 is None:
            return
        v = ask_string(self.root, "Link", "https://… or #3 (jump to page 3):",
                       e0["video"] if e0["t"] == "video" else e0["link"])
        if v is not None:
            self.set_link(v.strip())

    def edit_selected(self):
        el = self.cur()
        if el is not None and el["t"] in TEXTY:
            self.edit_text(self.sel[0])

    def edit_text(self, idx):
        if not (0 <= idx < len(self.els)):
            return
        el = self.els[idx]
        if el["t"] not in TEXTY:
            return
        self.end_edit()
        self.sel = [idx]
        self.sync_ui()
        self.editing = el
        z = self.z
        px = max(1, int(round(el["size"] * z)))
        bgc = blend(el["fill"], self.page["bg"], el["op"]) if el["fill"] else self.page["bg"]
        t = tk.Text(self.cv, font=self.tkfont(el, px), bg=bgc, fg=el["color"], insertbackground=el["color"], relief="flat",
                    bd=0, wrap="word", highlightthickness=2, highlightbackground=C["accent"], highlightcolor=C["accent"],
                    padx=max(0, int(6 * z)), pady=0, undo=True, selectbackground=C["accent2"], selectforeground="#ffffff")
        t.tag_configure("al", justify=el["align"])
        t.insert("1.0", el["text"])
        t.tag_add("al", "1.0", "end")
        self.editor = t

        def relayout(ev=None):
            if self.editor is not t:
                return
            tmp = dict(el, text=t.get("1.0", "end-1c"))
            try:
                lines, lh, bt = text_layout(tmp)
                t.config(pady=max(0, int(bt * self.z) - 2))
                t.tag_add("al", "1.0", "end")
            except Exception:
                pass
        t.bind("<KeyRelease>", relayout)
        t.bind("<Escape>", lambda ev: (self.end_edit(), self.cv.focus_set(), "break")[2])
        t.bind("<FocusOut>", lambda ev: self.root.after(1, self.end_edit))
        self.redraw()
        relayout()
        t.focus_set()
        t.tag_add("sel", "1.0", "end-1c")

    def end_edit(self, commit=True):
        t = self.editor
        if t is None:
            return
        el = self.editing
        self.editor = None
        self.editing = None
        try:
            txt = t.get("1.0", "end-1c")
        except Exception:
            txt = None
        try:
            t.destroy()
        except Exception:
            pass
        if commit and txt is not None and el is not None and el["text"] != txt and any(el is x for x in self.els):
            self.push_undo()
            el["text"] = txt
        self.sync_ui()
        self.redraw()
        self.schedule_thumb()

    def on_key_up(self, e):
        if e.keysym == "space" and self.space:
            self.space = False
            self.cv.config(cursor="arrow")

    def on_key(self, e):
        fw = self.root.focus_get()
        if isinstance(fw, (tk.Entry, tk.Text)):
            return
        k, ctrl, shift = e.keysym, bool(e.state & 0x4), bool(e.state & 0x1)
        if k == "space":
            self.space = True
            self.cv.config(cursor="fleur")
            return
        if ctrl:
            kl = k.lower()
            if kl == "z":
                self.do_redo() if shift else self.do_undo()
            elif kl == "y":
                self.do_redo()
            elif kl == "s":
                self.save_doc(shift)
            elif kl == "o":
                self.open_doc()
            elif kl == "n":
                self.new_doc()
            elif kl == "d":
                self.dup_el()
            elif kl == "e":
                self.export_pptx() if shift else self.export_pdf()
            elif kl == "c":
                self.copy()
            elif kl == "x":
                self.cut()
            elif kl == "v":
                self.paste()
            elif kl == "a":
                self.select_all()
            elif kl == "m":
                self.add_page()
            elif kl == "0":
                self.fit()
            elif kl == "1":
                self.set_zoom(1.0)
            elif kl in ("plus", "equal"):
                self.zoom_by(1.2)
            elif kl == "minus":
                self.zoom_by(1 / 1.2)
            return
        if k in ("Delete", "BackSpace"):
            self.del_el()
        elif k == "F5":
            self.present()
        elif k in ("F2", "Return", "KP_Enter"):
            self.edit_selected()
        elif k == "Escape":
            self.sel = []
            self.sync_ui()
            self.redraw()
        elif k in ("Prior", "Next") and not self.sel:
            self.goto(self.pi + (-1 if k == "Prior" else 1))
        elif k in ("plus", "equal", "KP_Add"):
            self.zoom_by(1.2)
        elif k in ("minus", "KP_Subtract"):
            self.zoom_by(1 / 1.2)
        elif k in ("Left", "Right", "Up", "Down") and self.sel:
            self.coalesce("nudge")
            st = 10 if shift else 1
            dx = -st if k == "Left" else st if k == "Right" else 0
            dy = -st if k == "Up" else st if k == "Down" else 0
            for el in self.selels():
                el["x"] += dx
                el["y"] += dy
            self.sync_ui()
            self.redraw()
            self.schedule_thumb()

    def confirm_discard(self):
        if not self.dirty:
            return True
        r = ask_save_discard(self.root, "PDFS", "Save your changes?")
        if r is None:
            return False
        if r:
            return self.save_doc()
        return True

    def reset_view(self):
        self.pil_cache.clear()
        self.cache_main.d.clear()
        self.cache_thumb.d.clear()
        self.pi, self.sel = 0, []
        self.undo.clear()
        self.redo.clear()
        self.auto_fit = True
        self.panx = self.pany = 0.0
        self.end_edit(commit=False)
        self.title()
        self.refresh_thumbs()
        self.sync_ui()
        self.redraw()

    def new_doc(self):
        if not self.confirm_discard():
            return
        self.W, self.H = SIZES["16:9  (960×540)"]
        self.pages, self.assets = [self.blank_page()], {}
        self.path, self.dirty = None, False
        self.reset_view()

    def save_doc(self, as_new=False):
        self.end_edit()
        p = self.path
        if as_new or not p:
            p = filedialog.asksaveasfilename(parent=self.root, title="Save project", defaultextension=".pdfs",
                                             initialdir=self.last_dir, filetypes=[("PDFS project", "*.pdfs")])
        if not p:
            return False
        try:
            used = set()
            for pg in self.pages:
                if pg["bgimg"]:
                    used.add(pg["bgimg"])
                for el in pg["els"]:
                    if el["asset"]:
                        used.add(el["asset"])
            with open(p, "w") as f:
                json.dump(dict(W=self.W, H=self.H, pages=self.pages, assets={k: v for k, v in self.assets.items() if k in used}), f)
        except Exception as ex:
            self.status(f"Save error: {ex}", "err")
            return False
        self.last_dir = os.path.dirname(p)
        self.path, self.dirty = p, False
        self.title()
        self.status(f"Saved: {os.path.basename(p)}")
        self._save_settings()
        return True

    def open_doc(self):
        if not self.confirm_discard():
            return
        p = filedialog.askopenfilename(parent=self.root, title="Open project", initialdir=self.last_dir,
                                       filetypes=[("PDFS project", "*.pdfs"), ("All files", "*.*")])
        if not p:
            return
        try:
            with open(p) as f:
                d = json.load(f)
            base = new_el("text", self.fams[0])
            pages = d["pages"]
            for pg in pages:
                pg["els"] = [dict(base, **el) for el in pg["els"]]
                for k, v in self.blank_page().items():
                    pg.setdefault(k, v)
            self.W, self.H, self.pages, self.assets = d["W"], d["H"], pages, d.get("assets", {})
        except Exception as ex:
            self.status(f"Open error: {ex}", "err")
            return
        self.last_dir = os.path.dirname(p)
        self.path, self.dirty = p, False
        self.reset_view()
        self.status(f"Opened: {os.path.basename(p)}")

    def export_pdf(self):
        self.end_edit()
        p = filedialog.asksaveasfilename(parent=self.root, title="Export PDF", defaultextension=".pdf",
                                         initialdir=self.last_dir,
                                         initialfile=(Path(self.path).stem if self.path else "presentation") + ".pdf",
                                         filetypes=[("PDF", "*.pdf")])
        if not p:
            return
        try:
            self.build_pdf(p)
        except Exception as ex:
            traceback.print_exc()
            self.status(f"PDF error: {ex}", "err")
            show_info(self.root, "PDF error", str(ex), accent="err")
            return
        self.last_dir = os.path.dirname(p)
        self.status(f"PDF ready: {os.path.basename(p)}  ({len(self.pages)} pages)")
        if ask_yes_no(self.root, "PDF ready", f"{os.path.basename(p)} was created.\nOpen it now?", yes="Open", no="Later"):
            try:
                webbrowser.open(Path(p).resolve().as_uri())
            except Exception:
                pass

    def export_pptx(self):
        self.end_edit()
        try:
            import pptx
        except ImportError:
            show_info(self.root, "PowerPoint export", "PowerPoint export needs python-pptx:\n\n    pip install python-pptx",
                      accent="warn")
            return
        p = filedialog.asksaveasfilename(parent=self.root, title="Export PowerPoint", defaultextension=".pptx",
                                         initialdir=self.last_dir,
                                         initialfile=(Path(self.path).stem if self.path else "presentation") + ".pptx",
                                         filetypes=[("PowerPoint", "*.pptx")])
        if not p:
            return
        try:
            self.build_pptx(p)
        except Exception as ex:
            traceback.print_exc()
            self.status(f"PPTX error: {ex}", "err")
            show_info(self.root, "PowerPoint error", str(ex), accent="err")
            return
        self.last_dir = os.path.dirname(p)
        self.status(f"PowerPoint ready: {os.path.basename(p)}  ({len(self.pages)} slides)")
        if ask_yes_no(self.root, "PowerPoint ready", f"{os.path.basename(p)} was created.\nOpen it now?", yes="Open", no="Later"):
            try:
                if sys.platform.startswith("win"):
                    os.startfile(p)
                else:
                    import subprocess
                    subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", p])
            except Exception:
                pass

    def build_pptx(self, path):
        from pptx import Presentation
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
        from pptx.oxml.ns import qn
        from lxml import etree

        W, H = self.W, self.H
        prs = Presentation()
        prs.slide_width, prs.slide_height = Pt(W), Pt(H)
        blank = prs.slide_layouts[6]

        def rgb(h):
            h = (h or "#000000").lstrip("#")
            if len(h) == 3:
                h = "".join(ch * 2 for ch in h)
            return RGBColor.from_string(h[:6].upper())

        def alpha(shape, op):
            if op >= 0.999:
                return
            for clr in shape._element.spPr.xpath(".//a:srgbClr"):
                for old in clr.findall(qn("a:alpha")):
                    clr.remove(old)
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(max(0.0, min(1.0, op)) * 100000)))

        def png(im):
            buf = io.BytesIO()
            im.save(buf, "PNG")
            buf.seek(0)
            return buf

        def add_el(s, el):
            t, op = el["t"], el["op"]
            x0, y0, x1, y1 = self.bbox_of([el])
            L, T, Wd, Ht = Pt(x0), Pt(y0), Pt(max(x1 - x0, 0.5)), Pt(max(y1 - y0, 0.5))
            if t == "line":
                c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(el["x"]), Pt(el["y"]),
                                           Pt(el["x"] + el["w"]), Pt(el["y"] + el["h"]))
                c.line.color.rgb = rgb(el["line"] or "#888888")
                c.line.width = Pt(max(0.25, el["lw"]))
                alpha(c, op)
                return []
            if t == "image":
                im = self.pil(el["asset"], op)
                return [s.shapes.add_picture(png(im), L, T, Wd, Ht)] if im is not None else []
            if t == "video":
                im = self.pil(el["asset"], op) if el["asset"] else None
                if im is not None:
                    base = s.shapes.add_picture(png(im), L, T, Wd, Ht)
                else:
                    base = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, T, Wd, Ht)
                    base.fill.solid()
                    base.fill.fore_color.rgb = rgb("#101010")
                    base.line.fill.background()
                    base.shadow.inherit = False
                r = min(x1 - x0, y1 - y0) * 0.14
                cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
                circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Pt(cx - r), Pt(cy - r), Pt(2 * r), Pt(2 * r))
                circ.fill.solid()
                circ.fill.fore_color.rgb = rgb("#ffffff")
                circ.line.fill.background()
                circ.shadow.inherit = False
                alpha(circ, 0.88 * op)
                tcx = cx + r * 0.125
                tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Pt(tcx - r * 0.55), Pt(cy - r * 0.475),
                                         Pt(r * 1.1), Pt(r * 0.95))
                tri.rotation = 90
                tri.fill.solid()
                tri.fill.fore_color.rgb = rgb("#222222")
                tri.line.fill.background()
                tri.shadow.inherit = False
                return [base, circ, tri]
            if t == "text":
                shp = s.shapes.add_textbox(L, T, Wd, Ht)
            else:
                kind = MSO_SHAPE.OVAL if t == "ellipse" else (
                    MSO_SHAPE.ROUNDED_RECTANGLE if el["r"] > 0 else MSO_SHAPE.RECTANGLE)
                shp = s.shapes.add_shape(kind, L, T, Wd, Ht)
                if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                    shp.adjustments[0] = max(0.0, min(0.5, el["r"] / max(1.0, min(x1 - x0, y1 - y0))))
            shp.shadow.inherit = False
            if el["fill"]:
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(el["fill"])
            else:
                shp.fill.background()
            if el["line"]:
                shp.line.color.rgb = rgb(el["line"])
                shp.line.width = Pt(max(0.25, el["lw"]))
            else:
                shp.line.fill.background()
            alpha(shp, op)
            if el["text"]:
                tf = shp.text_frame
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(6)
                va = el.get("va", "auto")
                if va == "auto":
                    va = "top" if t == "text" else "middle"
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                                      "bottom": MSO_ANCHOR.BOTTOM}[va]
                al = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[el["align"]]
                for k, ln in enumerate(el["text"].split("\n")):
                    para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    para.alignment = al
                    para.line_spacing = 1.0
                    para.font.size = Pt(el["size"])
                    run = para.add_run()
                    run.text = ln
                    f = run.font
                    f.name = el["font"]
                    f.size = Pt(el["size"])
                    f.bold = bool(el["b"])
                    f.italic = bool(el["i"])
                    f.color.rgb = rgb(el["color"])
            return [shp]

        slides, links = [], []
        for pg in self.pages:
            s = prs.slides.add_slide(blank)
            slides.append(s)
            bgim = self.pil(pg["bgimg"]) if pg["bgimg"] else None
            if bgim is not None:
                s.shapes.add_picture(png(bgim), 0, 0, Pt(W), Pt(H))
            else:
                f = s.background.fill
                done = False
                if pg["grad"]:
                    try:
                        f.gradient()
                        f.gradient_angle = 270
                        f.gradient_stops[0].color.rgb = rgb(pg["bg"])
                        f.gradient_stops[1].color.rgb = rgb(pg["bg2"])
                        done = True
                    except Exception:
                        done = False
                if not done:
                    f.solid()
                    f.fore_color.rgb = rgb(pg["bg"])
            for el in pg["els"]:
                shapes = add_el(s, el)
                tg = (el["video"] if el["t"] == "video" else el["link"]).strip()
                if shapes and tg and tg != "https://":
                    links.append((shapes, tg))
        for shapes, tg in links:
            if tg.startswith("#") and tg[1:].isdigit():
                n = int(tg[1:])
                if not (1 <= n <= len(slides)):
                    continue
            elif "://" not in tg and not tg.startswith("mailto:"):
                tg = "https://" + tg
            for shp in shapes:
                try:
                    if tg.startswith("#"):
                        shp.click_action.target_slide = slides[int(tg[1:]) - 1]
                    else:
                        shp.click_action.hyperlink.address = tg
                except Exception:
                    pass
        prs.save(path)

    def build_pdf(self, path):
        W, H = self.W, self.H
        c = rl_canvas.Canvas(path, pagesize=(W, H))
        c.setTitle(Path(self.path).stem if self.path else "PDFS")
        for pi, pg in enumerate(self.pages):
            c.bookmarkPage(f"p{pi + 1}")
            if pg["bgimg"] and self.pil(pg["bgimg"]):
                c.drawImage(ImageReader(self.pil(pg["bgimg"])), 0, 0, W, H)
            elif pg["grad"]:
                n = 64
                for i in range(n):
                    c.setFillColor(HexColor(lerp_hex(pg["bg"], pg["bg2"], i / (n - 1))))
                    c.rect(0, H - H * (i + 1) / n - 0.5, W, H / n + 1, stroke=0, fill=1)
            else:
                c.setFillColor(HexColor(pg["bg"]))
                c.rect(0, 0, W, H, stroke=0, fill=1)
            for el in pg["els"]:
                self.pdf_el(c, el, H)
            c.showPage()
        c.save()

    def pdf_el(self, c, el, H):
        t = el["t"]
        bx0, by0, bx1, by1 = self.bbox_of([el])
        x, w = bx0, bx1 - bx0
        h = by1 - by0
        y = H - by1
        c.saveState()
        try:
            c.setFillAlpha(el["op"])
            c.setStrokeAlpha(el["op"])
        except Exception:
            pass
        if t == "line":
            c.setStrokeColor(HexColor(el["line"] or "#888888"))
            c.setLineWidth(el["lw"])
            c.setLineCap(1)
            c.line(el["x"], H - el["y"], el["x"] + el["w"], H - el["y"] - el["h"])
        elif t == "image":
            im = self.pil(el["asset"])
            if im:
                c.drawImage(ImageReader(im), x, y, w, h, mask="auto")
        elif t == "video":
            im = self.pil(el["asset"]) if el["asset"] else None
            if im:
                c.drawImage(ImageReader(im), x, y, w, h, mask="auto")
            else:
                c.setFillColor(HexColor("#101010"))
                c.rect(x, y, w, h, stroke=0, fill=1)
            cx, cy, r = x + w / 2, y + h / 2, min(w, h) * 0.14
            c.setFillColor(HexColor("#ffffff"))
            c.setFillAlpha(0.88 * el["op"])
            c.circle(cx, cy, r, stroke=0, fill=1)
            c.setFillColor(HexColor("#222222"))
            pth = c.beginPath()
            pth.moveTo(cx - r * 0.35, cy + r * 0.55)
            pth.lineTo(cx - r * 0.35, cy - r * 0.55)
            pth.lineTo(cx + r * 0.6, cy)
            pth.close()
            c.drawPath(pth, stroke=0, fill=1)
        else:
            if el["fill"]:
                c.setFillColor(HexColor(el["fill"]))
            if el["line"]:
                c.setStrokeColor(HexColor(el["line"]))
                c.setLineWidth(el["lw"])
            fl, st = 1 if el["fill"] else 0, 1 if el["line"] else 0
            if fl or st:
                if t == "ellipse":
                    c.ellipse(x, y, x + w, y + h, stroke=st, fill=fl)
                elif el["r"] > 0:
                    c.roundRect(x, y, w, h, min(el["r"], w / 2, h / 2), stroke=st, fill=fl)
                else:
                    c.rect(x, y, w, h, stroke=st, fill=fl)
            if t in TEXTY and el["text"]:
                self.pdf_text(c, el, H)
        c.restoreState()
        target = el["video"] if t == "video" else el["link"]
        if target and target.strip() not in ("", "https://"):
            tg = target.strip()
            rect = (x, y, x + w, y + h)
            if tg.startswith("#") and tg[1:].isdigit():
                n = int(tg[1:])
                if 1 <= n <= len(self.pages):
                    c.linkAbsolute("", f"p{n}", rect, Border="[0 0 0]")
            else:
                if "://" not in tg and not tg.startswith("mailto:"):
                    tg = "https://" + tg
                c.linkURL(tg, rect, relative=0, thickness=0)

    def pdf_text(self, c, el, H):
        fname = pdf_font(el["font"], el["b"], el["i"])
        size, pad = el["size"], 6
        lines, lh, bt = text_layout(el)
        c.setFillColor(HexColor(el["color"]))
        c.setFont(fname, size)
        for k, ln in enumerate(lines):
            if not ln:
                continue
            by = H - (el["y"] + bt + k * lh + 0.88 * size)
            if el["align"] == "center":
                c.drawCentredString(el["x"] + el["w"] / 2, by, ln)
            elif el["align"] == "right":
                c.drawRightString(el["x"] + el["w"] - pad, by, ln)
            else:
                c.drawString(el["x"] + pad, by, ln)

    def present(self):
        self.end_edit()
        win = tk.Toplevel(self.root)
        win.configure(bg="#000000")
        win.attributes("-fullscreen", True)
        cvs = tk.Canvas(win, bg="#000000", highlightthickness=0)
        cvs.pack(fill="both", expand=True)
        st = dict(i=self.pi, z=1.0, ox=0, oy=0)
        cache = ImgCache(30)

        def show():
            cvs.delete("all")
            cache.trim()
            sw, sh = cvs.winfo_width(), cvs.winfo_height()
            z = min(sw / self.W, sh / self.H)
            ox, oy = (sw - self.W * z) / 2, (sh - self.H * z) / 2
            st.update(z=z, ox=ox, oy=oy)
            self.paint_page(cvs, self.pages[st["i"]], z, ox, oy, cache)
            self.mask_outside(cvs, ox, oy, z, sw, sh, "#000000")
            cvs.create_text(sw - 12, sh - 10, anchor="se", text=f"{st['i'] + 1}/{len(self.pages)}", fill="#444444", font=FONT)

        def nxt(e=None):
            st["i"] = min(len(self.pages) - 1, st["i"] + 1)
            show()

        def prv(e=None):
            st["i"] = max(0, st["i"] - 1)
            show()

        def close(e=None):
            win.destroy()
            self.cv.focus_set()

        def click(e):
            px, py = (e.x - st["ox"]) / st["z"], (e.y - st["oy"]) / st["z"]
            for el in reversed(self.pages[st["i"]]["els"]):
                x0, y0, x1, y1 = self.bbox_of([el])
                if x0 <= px <= x1 and y0 <= py <= y1:
                    tg = (el["video"] if el["t"] == "video" else el["link"]).strip()
                    if tg.startswith("#") and tg[1:].isdigit():
                        st["i"] = max(0, min(len(self.pages) - 1, int(tg[1:]) - 1))
                        show()
                        return
                    if tg and tg != "https://":
                        webbrowser.open(tg if "://" in tg else "https://" + tg)
                        return
            nxt()
        for k in ("<Right>", "<space>", "<Down>", "<Next>", "<Return>"):
            win.bind(k, nxt)
        for k in ("<Left>", "<Up>", "<Prior>", "<BackSpace>"):
            win.bind(k, prv)
        win.bind("<Escape>", close)
        cvs.bind("<Button-1>", click)
        cvs.bind("<Button-3>", prv)
        win.focus_force()
        win.after(200, show)

    def show_shortcuts(self):
        w = ToolWindow(self, "Keyboard Shortcuts", 640, 580, "⌨")
        rows = [("Ctrl+N / O / S", "New / open / save (Shift = Save As)"), ("Ctrl+E", "Export PDF"), ("Ctrl+Shift+E", "Export PowerPoint (.pptx)"),
                ("F5", "Present (arrows, Esc quits; links are clickable)"), ("Ctrl+Z / Y", "Undo / redo"),
                ("Ctrl+C / X / V", "Copy / cut / paste"), ("Ctrl+D", "Duplicate"),
                ("Ctrl+A", "Select all"), ("Del", "Delete"), ("Ctrl+M", "New page"),
                ("F2 / Enter / double-click", "Edit text in place (Esc = done)"),
                ("Shift + click", "Extend selection"), ("Drag on empty area", "Selection rectangle"),
                ("Shift + drag", "Axis-locked move / proportional resize"),
                ("Arrows (Shift = 10)", "Nudge"), ("Mouse wheel", "Zoom at cursor"),
                ("Middle button / Space+drag", "Pan the view"), ("Ctrl+0 / Ctrl+1", "Fit / 100%"),
                ("PgUp / PgDn", "Previous / next page"), ("Right-click", "Context menu"),
                ("Drag a thumbnail", "Reorder pages"), ("F1", "This window")]
        sf = ScrollFrame(w.body, bg=C["bg"])
        sf.pack(fill="both", expand=True)
        for a, b in rows:
            f = tk.Frame(sf.inner, bg=C["bg"])
            f.pack(fill="x", padx=12, pady=2)
            tk.Label(f, text=a, bg=C["bg"], fg=C["accent"], font=FONTB, width=28, anchor="w").pack(side="left")
            tk.Label(f, text=b, bg=C["bg"], fg=C["text"], font=FONT, anchor="w", justify="left").pack(side="left")
        sf.bind_tree(sf.inner)

    def show_about(self):
        w = ToolWindow(self, "About", 460, 260, "ℹ")
        tk.Label(w.body, text="PDFS", bg=C["bg"], fg=C["accent"], font=("Consolas", 18, "bold")).pack(pady=(26, 4))
        tk.Label(w.body, text="PowerPoint-like PDF editor\ntext · shapes · images · video links · page jumps\n\n"
                              "HexaMix / Flender style", bg=C["bg"], fg=C["text"], font=FONT, justify="center").pack()
        tk.Label(w.body, text="Python · tkinter · Pillow · reportlab", bg=C["bg"], fg=C["text3"], font=("Consolas", 8)).pack(pady=14)

    def on_close(self):
        self.end_edit()
        if self.confirm_discard():
            self._save_settings()
            self.root.destroy()


if __name__ == "__main__":
    App()
